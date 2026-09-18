"""Private Origen stormwater-hotspot ingestion and model concordance.

The customer workbooks remain outside the repository.  This module converts
them into a small, auditable serving bundle and deliberately treats hotspot
locations as static/weak evidence rather than event water-depth truth.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

SCHEMA_VERSION = "gwm.abu_dhabi_flood.origen_hotspots.v1"
DEFAULT_BUNDLE_ROOT = (
    Path.home() / ".local/share/gisdataagent/private/abu_dhabi_stormwater/origen_hotspots_batch1"
)

SOURCE_SPECS = (
    {
        "municipality": "ADM",
        "filename_contains": "ADM Stormwater Hotspot Analysis",
        "inventories": (
            {
                "inventory": "current",
                "inventory_version": "20260909-v4.4-current",
                "sheet": "ADM List",
                "header_row": 2,
            },
            {
                "inventory": "history",
                "inventory_version": "20260909-v4.4-previous-255",
                "sheet": "Sheet2",
                "header_row": 4,
            },
        ),
    },
    {
        "municipality": "AAM",
        "filename_contains": "AAM_Stormwater_Hotspot_Analysis",
        "inventories": (
            {
                "inventory": "current",
                "inventory_version": "20260909-v4.2-current",
                "sheet": "Analysis",
                "header_row": 3,
            },
        ),
    },
    {
        "municipality": "DRM",
        "filename_contains": "DRM_Stormwater_Hotspot_Analysis",
        "inventories": (
            {
                "inventory": "current",
                "inventory_version": "20260909-v4.2-current",
                "sheet": "Analysis",
                "header_row": 3,
            },
        ),
    },
)

FIELD_ALIASES = {
    "#": "hotspot_id",
    "hotspot_id": "hotspot_id",
    "hotspot_area": "hotspot_area",
    "hotspot_location": "hotspot_location",
    "latitude": "latitude",
    "longitude": "longitude",
    "criticality": "criticality",
    "developer_private_area": "developer_private_area",
    "developer_name": "developer_name",
    "network_available": "network_available",
    "root_cause": "root_cause",
    "root_cause_details": "root_cause_details",
    "idd_interventoin_planned": "idd_intervention_planned",
    "idd_intervention_planned": "idd_intervention_planned",
    "idd_design_category": "idd_design_category",
    "design_solution": "design_solution",
    "type_of_intervention": "intervention_type",
    "additional_budget_requirment": "additional_budget_required",
    "additional_budget_requested": "additional_budget_requested",
    "hotspot_network_rain_design_event_current": "current_design_event",
    "hotspot_network_design_technical_capacity_current": "current_network_capacity",
    "hotspot_network_rain_design_event_future": "future_design_event",
    "hotspot_network_design_technical_capacity_future": "future_network_capacity",
    "flood_volume_capacity_after_intervention": "flood_volume_capacity_after_intervention",
    "associated_recovery_time_after_intervention": "recovery_time_after_intervention",
    "planned_idd_mobilization_date": "planned_mobilization_date",
    "planned_idd_completion_date": "planned_completion_date",
    "critical_infrastructure_area_phase1_solution_completion_date": "phase1_completion_date",
    "project_stage": "project_stage",
    "approval_from_idas_secured_yes_no": "idas_approval",
    "design_description": "design_description",
    "actual_mobilization_date": "actual_mobilization_date",
    "actual_completion_date": "actual_completion_date",
    "intervention_status": "intervention_status",
    "idd_owner_name": "idd_owner_name",
    "project": "project",
    "consultant": "consultant",
    "notes": "notes",
}

INTERVENTION_FIELDS = (
    "idd_intervention_planned",
    "idd_design_category",
    "design_solution",
    "intervention_type",
    "additional_budget_required",
    "additional_budget_requested",
    "current_design_event",
    "current_network_capacity",
    "future_design_event",
    "future_network_capacity",
    "flood_volume_capacity_after_intervention",
    "recovery_time_after_intervention",
    "planned_mobilization_date",
    "planned_completion_date",
    "phase1_completion_date",
    "project_stage",
    "idas_approval",
    "design_description",
    "actual_mobilization_date",
    "actual_completion_date",
    "intervention_status",
    "idd_owner_name",
    "project",
    "consultant",
    "notes",
)

CRITICALITY_ORDER = {"Very High": 4, "High": 3, "Medium": 2, "Low": 1, "Unknown": 0}
DEPTH_THRESHOLDS_M = (0.01, 0.05)
EXCEL_ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
}


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _header_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("&", " and ")
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9#]+", "_", text)).strip("_")


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    if not text or text.lower() in {"n/a", "na", "none", "nil", "-"}:
        return None
    return text


def _parse_coordinate(value: Any, axis: str) -> tuple[float | None, list[str]]:
    issues: list[str] = []
    if isinstance(value, bool) or value is None:
        return None, [f"{axis}_missing"]
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).replace("\xa0", " ").strip()
        values = re.findall(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", text)
        if not values:
            return None, [f"{axis}_invalid"]
        number = float(values[0])
        if "°" in text or "º" in text:
            issues.append(f"{axis}_degree_symbol_removed")
        if len(values) > 1:
            issues.append(f"{axis}_embedded_coordinate_pair")
        if text != values[0]:
            issues.append(f"{axis}_text_normalized")
    limit = 90.0 if axis == "latitude" else 180.0
    if not math.isfinite(number) or abs(number) > limit:
        return None, [*issues, f"{axis}_out_of_range"]
    return number, issues


def _normalize_criticality(value: Any) -> str:
    compact = re.sub(r"[^a-z]", "", str(value or "").lower())
    return {
        "veryhigh": "Very High",
        "high": "High",
        "medium": "Medium",
        "low": "Low",
    }.get(compact, "Unknown")


def _normalize_yes_no(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"yes", "y", "true", "1"}:
        return "yes"
    if text in {"no", "n", "false", "0"}:
        return "no"
    return "unknown"


def _root_cause_category(value: Any, details: Any) -> str:
    text = f"{value or ''} {details or ''}".lower()
    if any(
        token in text
        for token in ("absence of drainage", "absense of drainage", "without network", "no network")
    ):
        return "absence_of_drainage_network"
    if any(token in text for token in ("inadequate design", "upgrade", "capacity")):
        return "insufficient_network_capacity"
    if any(token in text for token in ("block", "clogg", "maintenance", "cleaning")):
        return "blockage_or_maintenance"
    if any(
        token in text
        for token in ("slope", "level difference", "low point", "topograph", "depression")
    ):
        return "topography_or_low_point"
    return "other_or_unspecified"


def _record_from_rows(
    *,
    municipality: str,
    inventory: str,
    inventory_version: str,
    source_name: str,
    sheet_name: str,
    row_number: int,
    headers: Sequence[Any],
    values: Sequence[Any],
    formulas: Sequence[Any],
) -> tuple[dict[str, Any] | None, list[str]]:
    raw_values: dict[str, Any] = {}
    raw_formulas: dict[str, str] = {}
    normalized_columns: dict[str, Any] = {}
    seen: Counter[str] = Counter()
    for index, header in enumerate(headers):
        if header in (None, ""):
            continue
        label = str(header).strip()
        seen[label] += 1
        raw_label = label if seen[label] == 1 else f"{label}__{seen[label]}"
        value = values[index] if index < len(values) else None
        formula = formulas[index] if index < len(formulas) else None
        raw_values[raw_label] = _json_value(value)
        if isinstance(formula, str) and formula.startswith("="):
            raw_formulas[raw_label] = formula
        canonical = FIELD_ALIASES.get(_header_key(header))
        if canonical:
            normalized_columns[canonical] = value

    hotspot_id = _clean_text(normalized_columns.get("hotspot_id"))
    if not hotspot_id:
        return None, []
    latitude, lat_issues = _parse_coordinate(normalized_columns.get("latitude"), "latitude")
    longitude, lon_issues = _parse_coordinate(normalized_columns.get("longitude"), "longitude")
    issues = [*lat_issues, *lon_issues]
    if latitude is not None and longitude is not None and latitude > 40 and longitude < 40:
        latitude, longitude = longitude, latitude
        issues.append("coordinate_order_swapped")
    if (
        latitude is not None
        and longitude is not None
        and not (20 <= latitude <= 27 and 50 <= longitude <= 58)
    ):
        issues.append("coordinate_outside_uae_review_extent")

    criticality = _normalize_criticality(normalized_columns.get("criticality"))
    network_available = _normalize_yes_no(normalized_columns.get("network_available"))
    root_cause = _clean_text(normalized_columns.get("root_cause"))
    root_cause_details = _clean_text(normalized_columns.get("root_cause_details"))
    normalized = {
        "composite_id": f"{municipality}:{hotspot_id}:{inventory_version}",
        "municipality": municipality,
        "hotspot_id": hotspot_id,
        "inventory": inventory,
        "inventory_version": inventory_version,
        "hotspot_area": _clean_text(normalized_columns.get("hotspot_area")),
        "hotspot_location": _clean_text(normalized_columns.get("hotspot_location")),
        "latitude": latitude,
        "longitude": longitude,
        "criticality": criticality,
        "criticality_rank": CRITICALITY_ORDER[criticality],
        "network_available": network_available,
        "root_cause": root_cause,
        "root_cause_details": root_cause_details,
        "root_cause_category": _root_cause_category(root_cause, root_cause_details),
        "developer_private_area": _clean_text(normalized_columns.get("developer_private_area")),
        "developer_name": _clean_text(normalized_columns.get("developer_name")),
        "source_workbook": source_name,
        "source_sheet": sheet_name,
        "source_row": row_number,
        "coordinate_quality": "valid"
        if latitude is not None and longitude is not None
        else "invalid",
        "coordinate_normalizations": issues,
    }
    for field in INTERVENTION_FIELDS:
        normalized[field] = _json_value(normalized_columns.get(field))
    return {
        "normalized": normalized,
        "raw": raw_values,
        "source_formulas": raw_formulas,
    }, issues


def _resolve_workbooks(source_root: Path) -> list[tuple[dict[str, Any], Path]]:
    workbooks = [path for path in source_root.glob("*.xlsx") if not path.name.startswith("~$")]
    resolved: list[tuple[dict[str, Any], Path]] = []
    for spec in SOURCE_SPECS:
        matches = [path for path in workbooks if spec["filename_contains"] in path.name]
        if len(matches) != 1:
            raise ValueError(
                f"origen_workbook_resolution_failed:{spec['municipality']}:{len(matches)}"
            )
        resolved.append((spec, matches[0]))
    return resolved


def read_origen_workbooks(source_root: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Read the three customer workbooks without copying them into the repo."""

    import openpyxl

    records: list[dict[str, Any]] = []
    source_receipts: list[dict[str, Any]] = []
    issue_counts: Counter[str] = Counter()
    selected_formula_count = 0
    workbook_formula_count = 0
    external_formula_count = 0
    error_value_count = 0
    error_samples: list[dict[str, str]] = []

    for spec, path in _resolve_workbooks(source_root):
        formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        value_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        source_formula_count = 0
        source_external_formula_count = 0
        source_error_count = 0
        for formula_sheet in formula_book.worksheets:
            value_sheet = value_book[formula_sheet.title]
            for formula_row, value_row in zip(
                formula_sheet.iter_rows(), value_sheet.iter_rows(), strict=False
            ):
                for formula_cell, value_cell in zip(formula_row, value_row, strict=False):
                    formula_value = formula_cell.value
                    if isinstance(formula_value, str) and formula_value.startswith("="):
                        source_formula_count += 1
                        if re.search(
                            r"(?i)(?:https?://|sharepoint|\[[^\]]+\.(?:xlsx?|xlsm|xlsb)\])",
                            formula_value,
                        ):
                            source_external_formula_count += 1
                    cached_value = value_cell.value
                    if str(cached_value or formula_value) in EXCEL_ERROR_VALUES:
                        source_error_count += 1
                        if len(error_samples) < 50:
                            error_samples.append(
                                {
                                    "workbook": path.name,
                                    "sheet": formula_sheet.title,
                                    "cell": formula_cell.coordinate,
                                    "value": str(cached_value or formula_value),
                                }
                            )
        workbook_formula_count += source_formula_count
        external_formula_count += source_external_formula_count
        error_value_count += source_error_count
        source_record_count = 0
        for inventory_spec in spec["inventories"]:
            sheet_name = inventory_spec["sheet"]
            header_row = int(inventory_spec["header_row"])
            formula_sheet = formula_book[sheet_name]
            value_sheet = value_book[sheet_name]
            formula_rows = formula_sheet.iter_rows(min_row=header_row, values_only=True)
            value_rows = value_sheet.iter_rows(min_row=header_row, values_only=True)
            headers = next(value_rows)
            next(formula_rows)
            for row_number, (values, formulas) in enumerate(
                zip(value_rows, formula_rows, strict=False), start=header_row + 1
            ):
                record, issues = _record_from_rows(
                    municipality=str(spec["municipality"]),
                    inventory=str(inventory_spec["inventory"]),
                    inventory_version=str(inventory_spec["inventory_version"]),
                    source_name=path.name,
                    sheet_name=sheet_name,
                    row_number=row_number,
                    headers=headers,
                    values=values,
                    formulas=formulas,
                )
                if record is None:
                    continue
                records.append(record)
                source_record_count += 1
                issue_counts.update(issues)
                formulas_for_record = record["source_formulas"]
                selected_formula_count += len(formulas_for_record)
        source_receipts.append(
            {
                "municipality": spec["municipality"],
                "filename": path.name,
                "sha256": _sha256(path),
                "size_bytes": path.stat().st_size,
                "record_count": source_record_count,
                "formula_cell_count": source_formula_count,
                "external_formula_cell_count": source_external_formula_count,
                "cached_error_value_count": source_error_count,
                "external_link_relationship_count": len(
                    getattr(formula_book, "_external_links", ())
                ),
                "worksheets": [
                    {"name": sheet.title, "visibility": sheet.sheet_state}
                    for sheet in formula_book.worksheets
                ],
            }
        )
        formula_book.close()
        value_book.close()

    composite_ids = [record["normalized"]["composite_id"] for record in records]
    duplicate_ids = sorted(key for key, count in Counter(composite_ids).items() if count > 1)
    quality = {
        "schema": f"{SCHEMA_VERSION}.data_quality_receipt",
        "status": "passed_with_normalizations" if not duplicate_ids else "failed",
        "source_files": source_receipts,
        "record_count": len(records),
        "current_record_count": sum(
            record["normalized"]["inventory"] == "current" for record in records
        ),
        "history_record_count": sum(
            record["normalized"]["inventory"] == "history" for record in records
        ),
        "invalid_coordinate_count": sum(
            record["normalized"]["coordinate_quality"] != "valid" for record in records
        ),
        "coordinate_normalization_counts": dict(sorted(issue_counts.items())),
        "coordinate_normalized_record_count": sum(
            bool(record["normalized"]["coordinate_normalizations"]) for record in records
        ),
        "formula_cell_count_in_selected_rows": selected_formula_count,
        "formula_cell_count_workbook_wide": workbook_formula_count,
        "external_formula_cell_count_workbook_wide": external_formula_count,
        "cached_error_value_count_workbook_wide": error_value_count,
        "error_value_samples": error_samples,
        "duplicate_composite_ids": duplicate_ids,
        "claim_boundary": (
            "Hotspots are static inventory locations and weak spatial evidence; they are not "
            "event-specific flood extents, timings, recession observations, or water-depth truth."
        ),
    }
    return records, quality


def _iter_normalized(records: Iterable[dict[str, Any]], inventory: str | None = None):
    for record in records:
        normalized = record["normalized"]
        if inventory is None or normalized["inventory"] == inventory:
            yield normalized


def _attach_nearest_nodes(
    records: list[dict[str, Any]], node_path: Path | None, maximum_distance_m: float
) -> dict[str, Any]:
    if node_path is None:
        return {"status": "not_requested", "matched_count": 0}
    if not node_path.is_file():
        raise ValueError(f"hotspot_node_source_not_found:{node_path}")

    import geopandas as gpd
    import numpy as np
    from scipy.spatial import cKDTree

    nodes = gpd.read_file(node_path)
    if nodes.empty:
        raise ValueError("hotspot_node_source_empty")
    if nodes.crs is None:
        nodes = nodes.set_crs("EPSG:4326")
    projected = nodes.to_crs("EPSG:32640")
    coordinates = np.column_stack((projected.geometry.x, projected.geometry.y))
    tree = cKDTree(coordinates)
    current = [
        record["normalized"]
        for record in records
        if record["normalized"]["inventory"] == "current"
        and record["normalized"]["latitude"] is not None
        and record["normalized"]["longitude"] is not None
    ]
    points = gpd.GeoSeries.from_xy(
        [record["longitude"] for record in current],
        [record["latitude"] for record in current],
        crs="EPSG:4326",
    ).to_crs("EPSG:32640")
    query = np.column_stack((points.x, points.y))
    distances, indexes = tree.query(query, k=1)
    id_field = next(
        (field for field in ("node_id", "NODE_ID", "Node_ID", "id") if field in nodes.columns),
        None,
    )
    matched = 0
    by_municipality: Counter[str] = Counter()
    for record, distance, index in zip(current, distances, indexes, strict=True):
        accepted = float(distance) <= maximum_distance_m
        record["swmm_node_candidate"] = {
            "status": "candidate" if accepted else "outside_candidate_distance",
            "node_id": _json_value(nodes.iloc[int(index)][id_field])
            if accepted and id_field
            else None,
            "distance_m": round(float(distance), 3),
            "maximum_candidate_distance_m": maximum_distance_m,
            "binding_status": "candidate_only_not_engineering_confirmed",
        }
        if accepted:
            matched += 1
            by_municipality[record["municipality"]] += 1
    return {
        "status": "completed",
        "source_filename": node_path.name,
        "source_sha256": _sha256(node_path),
        "source_feature_count": len(nodes),
        "maximum_candidate_distance_m": maximum_distance_m,
        "matched_count": matched,
        "matched_by_municipality": dict(sorted(by_municipality.items())),
        "claim_boundary": (
            "Nearest-node links are spatial candidates and require asset-ID confirmation."
        ),
    }


def _attach_depth_concordance(
    records: list[dict[str, Any]], depth_results: Mapping[int, Path]
) -> dict[str, Any]:
    if not depth_results:
        return {
            "schema": f"{SCHEMA_VERSION}.spatial_concordance",
            "status": "not_requested",
            "return_periods": {},
        }

    from pyproj import Transformer
    from shapely.geometry import Point, shape
    from shapely.strtree import STRtree

    current = list(_iter_normalized(records, "current"))
    periods: dict[str, Any] = {}
    for return_period, source in sorted(depth_results.items()):
        if not source.is_file():
            raise ValueError(f"hotspot_depth_result_not_found:{source}")
        result_root = source.parent
        summary_path = result_root / "delivery_summary.json"
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        bounds = (summary.get("domain") or {}).get("bounds_epsg32640")
        if not isinstance(bounds, list) or len(bounds) != 4:
            raise ValueError(f"hotspot_depth_domain_bounds_missing:{summary_path}")
        collection = json.loads(source.read_text(encoding="utf-8"))
        features = collection.get("features") or []
        geometries = [shape(feature["geometry"]) for feature in features]
        depths = [
            float((feature.get("properties") or {}).get("maximum_depth_m") or 0.0)
            for feature in features
        ]
        tree = STRtree(geometries)
        transformer = Transformer.from_crs("EPSG:4326", "EPSG:32640", always_xy=True)
        municipality_counts: dict[str, Counter[str]] = {}
        misses: list[str] = []
        inside_count = 0
        hits = {threshold: 0 for threshold in DEPTH_THRESHOLDS_M}
        for record in current:
            municipality_counts.setdefault(record["municipality"], Counter())
            key = f"rp{return_period:03d}"
            latitude = record.get("latitude")
            longitude = record.get("longitude")
            if latitude is None or longitude is None:
                record.setdefault("depth_concordance", {})[key] = {
                    "domain_status": "invalid_coordinate",
                    "maximum_depth_m": None,
                }
                municipality_counts[record["municipality"]]["invalid_coordinate"] += 1
                continue
            x, y = transformer.transform(float(longitude), float(latitude))
            inside = bounds[0] <= x <= bounds[2] and bounds[1] <= y <= bounds[3]
            maximum_depth = 0.0
            if inside:
                point = Point(float(longitude), float(latitude))
                indexes = tree.query(point, predicate="intersects")
                if len(indexes):
                    maximum_depth = max(depths[int(index)] for index in indexes)
                inside_count += 1
                municipality_counts[record["municipality"]]["inside_domain"] += 1
                for threshold in DEPTH_THRESHOLDS_M:
                    if maximum_depth >= threshold:
                        hits[threshold] += 1
                        municipality_counts[record["municipality"]][f"hit_ge_{threshold:g}m"] += 1
                if maximum_depth < DEPTH_THRESHOLDS_M[0]:
                    misses.append(record["composite_id"])
            else:
                municipality_counts[record["municipality"]]["outside_domain"] += 1
            record.setdefault("depth_concordance", {})[key] = {
                "domain_status": "inside" if inside else "outside",
                "maximum_depth_m": maximum_depth if inside else None,
                "hit_ge_0_01m": inside and maximum_depth >= 0.01,
                "hit_ge_0_05m": inside and maximum_depth >= 0.05,
                "evidence_role": "static_hotspot_spatial_concordance_only",
            }
        periods[str(return_period)] = {
            "return_period_years": return_period,
            "source_filename": source.name,
            "source_sha256": _sha256(source),
            "domain_summary_sha256": _sha256(summary_path),
            "current_hotspot_count": len(current),
            "inside_domain_count": inside_count,
            "outside_domain_count": len(current) - inside_count,
            "hit_ge_0_01m_count": hits[0.01],
            "hit_ge_0_05m_count": hits[0.05],
            "miss_ge_0_01m_count": len(misses),
            "miss_ge_0_01m_composite_ids": misses,
            "by_municipality": {
                municipality: dict(sorted(counts.items()))
                for municipality, counts in sorted(municipality_counts.items())
            },
        }
    return {
        "schema": f"{SCHEMA_VERSION}.spatial_concordance",
        "status": "completed_weak_static_evidence",
        "return_periods": periods,
        "validation_gate": {
            "hotspot_spatial_concordance": "available",
            "event_water_depth_validation": "pending",
            "event_flood_extent_validation": "pending",
            "event_recession_time_validation": "pending",
        },
        "claim_boundary": (
            "Inventory points are compared with modeled depth grids only for spatial concordance. "
            "They do not provide event timing, observed depth, observed extent, or recession truth."
        ),
    }


def _geojson(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    features = []
    for record in records:
        latitude = record.get("latitude")
        longitude = record.get("longitude")
        if latitude is None or longitude is None:
            continue
        properties = {
            key: value for key, value in record.items() if key not in {"latitude", "longitude"}
        }
        node_candidate = properties.get("swmm_node_candidate")
        if isinstance(node_candidate, dict):
            properties["swmm_node_candidate_status"] = node_candidate.get("status")
            properties["swmm_node_candidate_id"] = node_candidate.get("node_id")
            properties["swmm_node_candidate_distance_m"] = node_candidate.get("distance_m")
        depth_concordance = properties.get("depth_concordance")
        if isinstance(depth_concordance, dict):
            for period_key, result in depth_concordance.items():
                if not isinstance(result, dict):
                    continue
                properties[f"{period_key}_domain_status"] = result.get("domain_status")
                properties[f"{period_key}_maximum_depth_m"] = result.get("maximum_depth_m")
                properties[f"{period_key}_hit_ge_0_01m"] = result.get("hit_ge_0_01m")
                properties[f"{period_key}_hit_ge_0_05m"] = result.get("hit_ge_0_05m")
        features.append(
            {
                "type": "Feature",
                "id": record["composite_id"],
                "geometry": {"type": "Point", "coordinates": [longitude, latitude]},
                "properties": properties,
            }
        )
    return {"type": "FeatureCollection", "features": features}


def _intervention_catalog(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    items = []
    for record in records:
        values = {
            field: record.get(field)
            for field in INTERVENTION_FIELDS
            if record.get(field) not in (None, "")
        }
        if not values:
            continue
        items.append(
            {
                "composite_id": record["composite_id"],
                "municipality": record["municipality"],
                "hotspot_id": record["hotspot_id"],
                "hotspot_location": record.get("hotspot_location"),
                "criticality": record["criticality"],
                "intervention": values,
                "engineering_effect": {
                    "status": "unquantified",
                    "modeled_depth_reduction_m": None,
                    "admission_requirement": (
                        "Bind an engineering design, hydraulic parameters, and an approved "
                        "scenario "
                        "before estimating flood-depth reduction."
                    ),
                },
            }
        )
    return {
        "schema": f"{SCHEMA_VERSION}.intervention_catalog",
        "status": "compiled_unquantified",
        "item_count": len(items),
        "items": items,
        "claim_boundary": (
            "Workbook intervention text is preserved; no text field is converted into a "
            "hydraulic benefit."
        ),
    }


def _has_intervention(record: Mapping[str, Any]) -> bool:
    return any(record.get(field) not in (None, "") for field in INTERVENTION_FIELDS)


def _intervention_state(record: Mapping[str, Any]) -> str:
    """Return a conservative inventory state without inferring hydraulic benefit."""

    status = re.sub(r"\s+", " ", str(record.get("intervention_status") or "")).strip().lower()
    stage = re.sub(r"\s+", " ", str(record.get("project_stage") or "")).strip().lower()
    if record.get("actual_completion_date") not in (None, "") or "completed" in status:
        return "completed"
    if "not started" in status:
        return "not_started"
    if any(
        token in f"{status} {stage}"
        for token in ("on track", "under construction", "implementation", "execution")
    ):
        return "active"
    return "recorded" if _has_intervention(record) else "none"


def _build_static_prior_features(
    records: list[dict[str, Any]],
    model_grid_path: Path | None,
    output_root: Path,
    *,
    influence_scale_m: float,
    cutoff_m: float,
) -> dict[str, Any]:
    """Rasterize Origen inventory attributes as prospective GWM covariates.

    The feature cube is aligned to the existing 250 m model grid.  It is kept
    separate from frozen models and confirmatory observations so that using it
    later requires a new training freeze and spatially blocked ablation.
    """

    if model_grid_path is None:
        return {
            "status": "not_requested",
            "admission": "prospective_training_only",
        }
    if influence_scale_m <= 0.0 or cutoff_m <= 0.0:
        raise ValueError("hotspot_static_prior_kernel_invalid")

    import numpy as np
    from pyproj import Transformer
    from scipy.spatial import cKDTree

    grid_path = model_grid_path.expanduser().resolve()
    if not grid_path.is_file():
        raise ValueError(f"hotspot_static_prior_grid_not_found:{grid_path}")
    with np.load(grid_path) as archive:
        x = np.asarray(archive["x"], dtype=np.float64)
        y = np.asarray(archive["y"], dtype=np.float64)
        land = np.asarray(archive["land_mask"], dtype=bool)
    rows, columns = land.shape
    if x.size == columns + 1:
        cell_x = 0.5 * (x[:-1] + x[1:])
    elif x.size == columns:
        cell_x = x
    else:
        raise ValueError("hotspot_static_prior_grid_x_invalid")
    if y.size == rows + 1:
        cell_y = 0.5 * (y[:-1] + y[1:])
    elif y.size == rows:
        cell_y = y
    else:
        raise ValueError("hotspot_static_prior_grid_y_invalid")
    grid_x, grid_y = np.meshgrid(cell_x, cell_y)
    cells = np.column_stack((grid_x.reshape(-1), grid_y.reshape(-1)))
    land_flat = land.reshape(-1)
    transformer = Transformer.from_crs("EPSG:4326", "EPSG:32640", always_xy=True)

    def projected(
        source: Iterable[Mapping[str, Any]],
    ) -> tuple[np.ndarray, list[Mapping[str, Any]]]:
        selected = [
            record
            for record in source
            if record.get("longitude") is not None and record.get("latitude") is not None
        ]
        if not selected:
            return np.empty((0, 2), dtype=np.float64), []
        eastings, northings = transformer.transform(
            [float(record["longitude"]) for record in selected],
            [float(record["latitude"]) for record in selected],
        )
        return np.column_stack((eastings, northings)).astype(np.float64), selected

    def maximum_influence(points: np.ndarray, weights: np.ndarray | None = None) -> np.ndarray:
        result = np.zeros(cells.shape[0], dtype=np.float32)
        if not len(points):
            return result
        point_weights = (
            np.ones(len(points), dtype=np.float32)
            if weights is None
            else np.asarray(weights, dtype=np.float32)
        )
        for start in range(0, len(points), 64):
            chunk = points[start : start + 64]
            difference = cells[:, None, :] - chunk[None, :, :]
            distance = np.sqrt(np.sum(difference * difference, axis=2))
            influence = np.exp(-distance / influence_scale_m).astype(np.float32)
            influence[distance > cutoff_m] = 0.0
            influence *= point_weights[start : start + 64][None, :]
            result = np.maximum(result, influence.max(axis=1))
        result[~land_flat] = 0.0
        return result

    current_records = list(_iter_normalized(records, "current"))
    history_records = list(_iter_normalized(records, "history"))
    current_points, current_valid = projected(current_records)
    history_points, history_valid = projected(history_records)
    land_tree = cKDTree(cells[land_flat])
    current_distance, _ = land_tree.query(current_points, k=1)
    history_distance, _ = land_tree.query(history_points, k=1)
    current_contributes = current_distance <= cutoff_m
    history_contributes = history_distance <= cutoff_m
    proximity = maximum_influence(current_points)
    density = np.zeros(cells.shape[0], dtype=np.float32)
    for start in range(0, len(current_points), 64):
        chunk = current_points[start : start + 64]
        difference = cells[:, None, :] - chunk[None, :, :]
        squared_distance = np.sum(difference * difference, axis=2)
        density += np.sum(squared_distance <= 1_000.0**2, axis=1, dtype=np.int32)
    density = np.clip(density / 4.0, 0.0, 1.0)
    density[~land_flat] = 0.0

    def selected_influence(predicate) -> np.ndarray:
        indexes = [index for index, record in enumerate(current_valid) if predicate(record)]
        return maximum_influence(current_points[indexes]) if indexes else np.zeros_like(proximity)

    criticality = maximum_influence(
        current_points,
        np.asarray(
            [float(record.get("criticality_rank") or 0.0) / 4.0 for record in current_valid],
            dtype=np.float32,
        ),
    )
    features: list[tuple[str, np.ndarray]] = [
        ("origen_current_hotspot_proximity_exp_1km", proximity),
        ("origen_current_hotspot_density_1km_capped_4", density),
        ("origen_criticality_weighted_proximity_exp_1km", criticality),
        (
            "origen_network_absence_proximity_exp_1km",
            selected_influence(lambda record: record.get("network_available") == "no"),
        ),
        ("origen_history_hotspot_proximity_exp_1km", maximum_influence(history_points)),
    ]
    for category in (
        "absence_of_drainage_network",
        "insufficient_network_capacity",
        "blockage_or_maintenance",
        "topography_or_low_point",
        "other_or_unspecified",
    ):
        features.append(
            (
                f"origen_root_cause_{category}_proximity_exp_1km",
                selected_influence(
                    lambda record, value=category: record.get("root_cause_category")
                    == value
                ),
            )
        )
    features.extend(
        [
            (
                "origen_intervention_recorded_proximity_exp_1km",
                selected_influence(_has_intervention),
            ),
            (
                "origen_intervention_not_started_proximity_exp_1km",
                selected_influence(lambda record: _intervention_state(record) == "not_started"),
            ),
            (
                "origen_intervention_active_proximity_exp_1km",
                selected_influence(lambda record: _intervention_state(record) == "active"),
            ),
            (
                "origen_intervention_completed_proximity_exp_1km",
                selected_influence(lambda record: _intervention_state(record) == "completed"),
            ),
            (
                "origen_swmm_node_candidate_proximity_exp_1km",
                selected_influence(
                    lambda record: isinstance(record.get("swmm_node_candidate"), dict)
                    and record["swmm_node_candidate"].get("status") == "candidate"
                ),
            ),
        ]
    )
    feature_names = [name for name, _ in features]
    cube = np.stack([values.reshape(land.shape) for _, values in features]).astype(np.float32)
    output_path = output_root / "gwm_static_prior_250m.npz"
    np.savez_compressed(
        output_path,
        features=cube,
        feature_names=np.asarray(feature_names),
        x=cell_x,
        y=cell_y,
        land_mask=land,
        epsg=np.asarray(32640, dtype=np.int32),
    )
    statistics = {
        name: {
            "minimum": float(np.min(values)),
            "maximum": float(np.max(values)),
            "mean_on_land": float(np.mean(values[land_flat])) if np.any(land_flat) else 0.0,
            "nonzero_land_cell_count": int(np.count_nonzero(values[land_flat])),
        }
        for name, values in features
    }
    zero_feature_names = [
        name for name, values in features if not np.any(values[land_flat] > 0.0)
    ]
    contributing_by_municipality = Counter(
        record["municipality"]
        for record, contributes in zip(current_valid, current_contributes, strict=True)
        if contributes
    )
    intervention_mask = np.asarray(
        [_has_intervention(record) for record in current_valid], dtype=bool
    )
    receipt = {
        "schema": f"{SCHEMA_VERSION}.static_prior.v1",
        "status": "ready_prospective_training_only",
        "artifact": {
            "filename": output_path.name,
            "sha256": _sha256(output_path),
            "size_bytes": output_path.stat().st_size,
        },
        "grid": {
            "source_filename": grid_path.name,
            "source_sha256": _sha256(grid_path),
            "epsg": 32640,
            "shape": [rows, columns],
            "cell_size_m": float(np.median(np.diff(cell_x))) if cell_x.size > 1 else None,
            "land_cell_count": int(np.count_nonzero(land)),
        },
        "kernel": {
            "type": "exponential_distance_decay",
            "influence_scale_m": influence_scale_m,
            "cutoff_m": cutoff_m,
            "density_radius_m": 1_000.0,
            "density_cap_count": 4,
        },
        "source_record_counts": {
            "current": len(current_records),
            "current_valid_coordinate": len(current_valid),
            "current_contributing_within_cutoff": int(np.count_nonzero(current_contributes)),
            "current_contributing_by_municipality": dict(
                sorted(contributing_by_municipality.items())
            ),
            "history": len(history_records),
            "history_valid_coordinate": len(history_valid),
            "history_contributing_within_cutoff": int(np.count_nonzero(history_contributes)),
            "intervention_recorded": int(np.count_nonzero(intervention_mask)),
            "intervention_contributing_within_cutoff": int(
                np.count_nonzero(intervention_mask & current_contributes)
            ),
        },
        "feature_names": feature_names,
        "active_feature_count": len(feature_names) - len(zero_feature_names),
        "zero_feature_names": zero_feature_names,
        "feature_statistics": statistics,
        "admission": {
            "allowed": "new_prospectively_frozen_gwm_training_or_scenario_conditioning",
            "required_validation": (
                "spatially_blocked_ablation_against_the_same_model_without_origen_features"
            ),
            "forbidden": [
                "retrofit_into_existing_frozen_models",
                "event_water_depth_label",
                "event_flood_extent_ground_truth",
            ],
        },
        "claim_boundary": (
            "The cube encodes proximity to a current static inventory. It does not encode "
            "event timing, observed depth, flood extent, intervention benefit, or recession truth."
        ),
    }
    _write_json(output_root / "gwm_static_prior_receipt.json", receipt)
    return receipt


def build_private_bundle(
    source_root: Path,
    output_root: Path,
    *,
    depth_results: Mapping[int, Path] | None = None,
    node_path: Path | None = None,
    maximum_node_distance_m: float = 500.0,
    model_grid_path: Path | None = None,
    hotspot_influence_scale_m: float = 1_000.0,
    hotspot_cutoff_m: float = 3_000.0,
) -> dict[str, Any]:
    """Build an auditable private bundle from Origen Batch1 workbooks."""

    import pandas as pd

    source_root = source_root.expanduser().resolve()
    output_root = output_root.expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    records, quality = read_origen_workbooks(source_root)
    network_join = _attach_nearest_nodes(records, node_path, maximum_node_distance_m)
    concordance = _attach_depth_concordance(records, depth_results or {})
    current = list(_iter_normalized(records, "current"))
    history = list(_iter_normalized(records, "history"))
    interventions = _intervention_catalog(current)
    static_prior = _build_static_prior_features(
        records,
        model_grid_path,
        output_root,
        influence_scale_m=hotspot_influence_scale_m,
        cutoff_m=hotspot_cutoff_m,
    )

    current_geojson = _geojson(current)
    history_geojson = _geojson(history)
    _write_json(output_root / "hotspots_current.geojson", current_geojson)
    _write_json(output_root / "hotspots_history.geojson", history_geojson)
    _write_json(output_root / "data_quality_receipt.json", quality)
    _write_json(output_root / "hotspot_spatial_concordance.json", concordance)
    _write_json(output_root / "intervention_catalog.json", interventions)
    with (output_root / "hotspot_records.jsonl").open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(
                json.dumps(record, ensure_ascii=False, sort_keys=True, allow_nan=False) + "\n"
            )

    frame_rows = []
    for record in records:
        normalized = dict(record["normalized"])
        # Workbook intervention columns mix cached numbers, dates, text, and
        # formulas between municipalities.  Keep their semantic content as
        # text in the columnar export so Arrow does not infer an unstable
        # numeric type from the first municipality.
        for field in INTERVENTION_FIELDS:
            value = normalized.get(field)
            normalized[field] = None if value is None else str(value)
        normalized["raw_values_json"] = json.dumps(
            record["raw"], ensure_ascii=False, sort_keys=True
        )
        normalized["source_formulas_json"] = json.dumps(
            record["source_formulas"], ensure_ascii=False, sort_keys=True
        )
        normalized["coordinate_normalizations_json"] = json.dumps(
            normalized.pop("coordinate_normalizations"), ensure_ascii=False
        )
        normalized["swmm_node_candidate_json"] = json.dumps(
            normalized.pop("swmm_node_candidate", None), ensure_ascii=False, sort_keys=True
        )
        normalized["depth_concordance_json"] = json.dumps(
            normalized.pop("depth_concordance", None), ensure_ascii=False, sort_keys=True
        )
        frame_rows.append(normalized)
    pd.DataFrame(frame_rows).to_parquet(output_root / "hotspots.parquet", index=False)

    criticality = Counter(record["criticality"] for record in current)
    municipalities = Counter(record["municipality"] for record in current)
    network = Counter(record["network_available"] for record in current)
    artifacts = {}
    for name in (
        "hotspots_current.geojson",
        "hotspots_history.geojson",
        "hotspots.parquet",
        "hotspot_records.jsonl",
        "data_quality_receipt.json",
        "hotspot_spatial_concordance.json",
        "intervention_catalog.json",
    ):
        path = output_root / name
        artifacts[name] = {"sha256": _sha256(path), "size_bytes": path.stat().st_size}
    if static_prior["status"] != "not_requested":
        for name in ("gwm_static_prior_250m.npz", "gwm_static_prior_receipt.json"):
            path = output_root / name
            artifacts[name] = {"sha256": _sha256(path), "size_bytes": path.stat().st_size}
    manifest = {
        "schema": f"{SCHEMA_VERSION}.bundle_manifest",
        "status": "ready_private_derived_bundle",
        "created_at_utc": datetime.now(UTC).isoformat().replace("+00:00", "Z"),
        "inventory": {
            "current_count": len(current),
            "history_count": len(history),
            "current_by_municipality": dict(sorted(municipalities.items())),
            "current_by_criticality": {
                key: criticality.get(key, 0)
                for key in ("Very High", "High", "Medium", "Low", "Unknown")
            },
            "current_by_network_availability": dict(sorted(network.items())),
        },
        "network_candidate_join": network_join,
        "spatial_concordance_status": concordance["status"],
        "intervention_item_count": interventions["item_count"],
        "gwm_static_prior": {
            "status": static_prior["status"],
            "feature_count": len(static_prior.get("feature_names", [])),
            "active_feature_count": static_prior.get("active_feature_count", 0),
            "zero_feature_names": static_prior.get("zero_feature_names", []),
            "source_record_counts": static_prior.get("source_record_counts", {}),
            "admission": static_prior.get("admission"),
            "claim_boundary": static_prior.get("claim_boundary"),
        },
        "artifacts": artifacts,
        "privacy": {
            "classification": "customer_private_derived",
            "raw_workbooks_copied": False,
            "repository_admission": "forbidden",
            "service_exposure": "authenticated_normalized_fields_only",
        },
        "gwm_admission": {
            "allowed_roles": ["auxiliary_static_feature", "spatial_prior", "external_diagnostic"],
            "forbidden_roles": ["event_water_depth_label", "event_flood_extent_ground_truth"],
            "frozen_confirmatory_model_use": "forbidden",
            "prospective_training_requirement": "new_freeze_and_spatially_blocked_ablation",
        },
    }
    _write_json(output_root / "manifest.json", manifest)
    return manifest


def augment_private_bundle_with_static_prior(
    output_root: Path,
    model_grid_path: Path,
    *,
    hotspot_influence_scale_m: float = 1_000.0,
    hotspot_cutoff_m: float = 3_000.0,
) -> dict[str, Any]:
    """Add a model-grid feature cube without rebuilding established evidence artifacts."""

    output_root = output_root.expanduser().resolve()
    records_path = output_root / "hotspot_records.jsonl"
    manifest_path = output_root / "manifest.json"
    if not records_path.is_file() or not manifest_path.is_file():
        raise ValueError("abu_dhabi_hotspot_bundle_not_available")
    records = []
    for line_number, line in enumerate(records_path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        record = json.loads(line)
        if not isinstance(record, dict) or not isinstance(record.get("normalized"), dict):
            raise ValueError(f"abu_dhabi_hotspot_record_invalid:{line_number}")
        records.append(record)
    static_prior = _build_static_prior_features(
        records,
        model_grid_path,
        output_root,
        influence_scale_m=hotspot_influence_scale_m,
        cutoff_m=hotspot_cutoff_m,
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(manifest, dict):
        raise ValueError("abu_dhabi_hotspot_manifest_invalid")
    manifest["gwm_static_prior"] = {
        "status": static_prior["status"],
        "feature_count": len(static_prior.get("feature_names", [])),
        "active_feature_count": static_prior.get("active_feature_count", 0),
        "zero_feature_names": static_prior.get("zero_feature_names", []),
        "source_record_counts": static_prior.get("source_record_counts", {}),
        "admission": static_prior.get("admission"),
        "claim_boundary": static_prior.get("claim_boundary"),
    }
    manifest.setdefault("gwm_admission", {}).update(
        {
            "frozen_confirmatory_model_use": "forbidden",
            "prospective_training_requirement": "new_freeze_and_spatially_blocked_ablation",
        }
    )
    artifacts = manifest.setdefault("artifacts", {})
    for name in ("gwm_static_prior_250m.npz", "gwm_static_prior_receipt.json"):
        path = output_root / name
        artifacts[name] = {"sha256": _sha256(path), "size_bytes": path.stat().st_size}
    manifest["updated_at_utc"] = datetime.now(UTC).isoformat().replace("+00:00", "Z")
    _write_json(manifest_path, manifest)
    return manifest


def configured_bundle_root() -> Path:
    value = os.environ.get("ABU_DHABI_HOTSPOT_BUNDLE_ROOT", "").strip()
    return Path(value).expanduser().resolve() if value else DEFAULT_BUNDLE_ROOT.resolve()


def _read_bundle_json(filename: str) -> dict[str, Any]:
    path = configured_bundle_root() / filename
    if not path.is_file():
        raise ValueError("abu_dhabi_hotspot_bundle_not_available")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("abu_dhabi_hotspot_bundle_invalid")
    return payload


def hotspot_catalog_payload() -> dict[str, Any]:
    manifest = _read_bundle_json("manifest.json")
    quality = _read_bundle_json("data_quality_receipt.json")
    concordance = _read_bundle_json("hotspot_spatial_concordance.json")
    interventions = _read_bundle_json("intervention_catalog.json")
    return {
        "schema": f"{SCHEMA_VERSION}.catalog",
        "status": manifest.get("status"),
        "inventory": manifest.get("inventory"),
        "network_candidate_join": manifest.get("network_candidate_join"),
        "quality": {
            "status": quality.get("status"),
            "invalid_coordinate_count": quality.get("invalid_coordinate_count"),
            "coordinate_normalization_counts": quality.get("coordinate_normalization_counts"),
        },
        "spatial_concordance": concordance,
        "intervention_catalog": {
            "status": interventions.get("status"),
            "item_count": interventions.get("item_count"),
            "claim_boundary": interventions.get("claim_boundary"),
        },
        "gwm_static_prior": manifest.get("gwm_static_prior"),
        "privacy": manifest.get("privacy"),
        "gwm_admission": manifest.get("gwm_admission"),
    }


def hotspot_geojson_payload(inventory: str = "current") -> dict[str, Any]:
    if inventory not in {"current", "history"}:
        raise ValueError("abu_dhabi_hotspot_inventory_invalid")
    return _read_bundle_json(f"hotspots_{inventory}.geojson")
