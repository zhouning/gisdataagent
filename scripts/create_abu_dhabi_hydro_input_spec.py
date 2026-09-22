#!/usr/bin/env python3
"""Build the client-facing Abu Dhabi hydro-model input and ETL specification."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.create_abu_dhabi_flood_data_requirements import (  # noqa: E402
    CURRENT,
    HYDRO,
    PRIORITY,
    QUALITY,
    TEMPLATES,
)

OUT_DIR = ROOT / "docs" / "deliverables"
OUT_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT_DIR / "阿布扎比暴雨内涝模型_一二维水动力模型客户数据输入与ETL规范_V1.0.docx"
XLSX_PATH = OUT_DIR / "阿布扎比暴雨内涝模型_一二维水动力模型输入字段矩阵_V1.0.xlsx"
JSON_PATH = OUT_DIR / "阿布扎比暴雨内涝模型_一二维水动力模型输入字段矩阵_V1.0.json"

NAVY = "15324B"
BLUE = "1F6FEB"
CYAN = "0B7285"
LIGHT_BLUE = "EAF3FF"
LIGHT_CYAN = "E9F7F8"
LIGHT_ORANGE = "FFF4E5"
LIGHT_RED = "FDECEC"
LIGHT_GREEN = "EAF7EE"
GRID = "C9D5E2"
TEXT = "1D2A39"
MUTED = "5C6B7A"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "4") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 70, start: int = 80, bottom: int = 70, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, value: object, *, bold: bool = False, color: str = TEXT, size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run("" if value is None else str(value))
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_border(cell)
    set_cell_margins(cell)


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[object]], widths: Sequence[float], *, font_size: float = 7.8, header_fill: str = NAVY):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, (cell, heading) in enumerate(zip(header.cells, headers)):
        cell.width = Inches(widths[idx])
        set_cell_text(cell, heading, bold=True, color=WHITE, size=font_size)
        set_cell_shading(cell, header_fill)
    for row_idx, values in enumerate(rows, start=1):
        row = table.add_row()
        fill = "FFFFFF" if row_idx % 2 else "F7FAFC"
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[idx])
            set_cell_text(cell, value, size=font_size)
            set_cell_shading(cell, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.2)
    run.font.color.rgb = RGBColor.from_string(TEXT)


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_ORANGE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="E2B66C", size="6")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run2 = p2.add_run(body)
    run2.font.name = "Arial"
    run2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run2.font.size = Pt(9.2)
    run2.font.color.rgb = RGBColor.from_string(TEXT)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor.from_string(NAVY if level <= 2 else CYAN)
    return p


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(NAVY)
        text = text[len(bold_prefix):]
    r2 = p.add_run(text)
    for run in p.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9.5)
        if not run.font.color.rgb:
            run.font.color.rgb = RGBColor.from_string(TEXT)


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    for style_name, size, color in [("Title", 22, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 12, CYAN), ("Heading 3", 10.5, CYAN)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Abu Dhabi Stormwater Model Input & ETL Specification · V1.0 · ")
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(footer)


def build_docx() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(35)
    title.paragraph_format.space_after = Pt(7)
    r = title.add_run("阿布扎比城市暴雨内涝世界模型 V1.1")
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("一维 / 二维水动力模型客户数据输入与 ETL 交付规范")
    sr.font.name = "Arial"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(15)
    sr.font.color.rgb = RGBColor.from_string(CYAN)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("版本：V1.0　日期：2026-09-22　范围：EPA SWMM 1D + ANUGA 2D + 一二维耦合")
    mr.font.name = "Arial"
    mr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    mr.font.size = Pt(9.5)
    mr.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph()

    add_callout(
        doc,
        "先给结论：客户原始数据不能直接作为模型输入。",
        "除非客户交付的文件已经符合本规范并通过质量门，否则客户交付的是源数据，不是可直接运行的模型输入。模型需要经过坐标、单位、时间、ID、拓扑、质量和版本治理后的原生输入。泵站数据尤其如此：只有静态泵资产清单不能驱动 SWMM，原始 SCADA 也不能直接写入 [PUMPS] / [CURVES] / [CONTROLS]，必须先完成 ETL、设备映射、曲线绑定、时间对齐和质量检查。",
        LIGHT_ORANGE,
    )

    add_heading(doc, "1. 文档目的与使用方式", 1)
    add_body(doc, "本文件把客户需要交付的数据，转换为可执行的数据契约：每一类数据应采用什么格式、包含哪些字段、如何组织记录、通过哪些质量检查，以及 ETL 后要生成什么 SWMM / ANUGA 输入对象。")
    add_body(doc, "本文件同时解释当前 V1.1 页面中的参数来源：CUSTOMER 表示已接入客户源数据，DERIVED 表示由 ETL 或模型适配器生成，ASSUMED / DEFAULT 表示当前为保证诊断性运行而使用的替代值。")

    add_heading(doc, "2. 三层数据契约：原始数据不等于模型输入", 1)
    add_table(
        doc,
        ["层级", "交付对象", "典型格式", "是否可直接运行", "必须完成的工作"],
        [
            ["L0 客户源数据", "管网 GDB、DTM、SCADA、潮位、降雨、观测、图纸和日志", "FileGDB / GeoPackage / PostGIS / GeoTIFF / CSV / Parquet / PDF", "否", "读取、解码、字段字典、来源和版本登记"],
            ["L1 标准化数据层", "统一 ID、CRS、垂直基准、UTC 时间、单位和代码域的表/图层", "GeoPackage / PostGIS / Parquet / COG / NetCDF", "否；用于质量检查", "拓扑修复、时间对齐、单位转换、质量码和缺失掩码"],
            ["L2 模型原生输入", "SWMM .inp、ANUGA mesh/domain/quantity/boundary、耦合映射和运行配置", ".inp / .msh or .tsh / GeoTIFF or NumPy / JSONL / Parquet", "通过质量门后才可运行", "编译模型、生成 manifest、哈希和 run receipt"],
        ],
        [1.0, 2.35, 2.0, 1.05, 2.45],
        font_size=8.2,
    )

    add_heading(doc, "3. 当前数据基础与尚未满足的核心输入", 1)
    add_body(doc, "当前项目已经具备管网空间骨架和客户 5 m DTM，可支持 ETL、拓扑诊断和研究型/诊断型模拟；这不表示管网水力属性、边界、泵站动作和校准观测已经达到工程入模条件。")
    current_rows = []
    for row in CURRENT:
        current_rows.append([row[0], row[1], row[3], row[4], row[6]])
    section = doc.add_section(start_type=1)
    set_landscape(section)
    add_table(doc, ["编号", "当前数据", "状态", "当前内容", "主要边界"], current_rows, [0.65, 1.45, 1.3, 4.2, 3.0], font_size=7.7)

    add_heading(doc, "4. 按数据类别的输入格式与结构要求", 1)
    add_body(doc, "下面 34 项是当前一维、二维和一二维耦合模型需要治理的主要数据项。P0 表示没有该数据时只能保持诊断性/研究型结果，不能完成工程校准或正式验收；P1 表示显著增益项。")
    hydro_groups = [
        ("4.1 一维网络、泵站与控制设施", HYDRO[:12]),
        ("4.2 二维地形、下垫面、汇水区与一二维交换", HYDRO[12:19]),
        ("4.3 降雨、潮位、运行时序、观测与事件包", HYDRO[19:32]),
        ("4.4 数据治理、版本与可追溯性", HYDRO[32:]),
    ]
    for heading, rows in hydro_groups:
        add_heading(doc, heading, 2)
        compact = []
        for row in rows:
            compact.append([
                row[0],
                row[1],
                f"{row[2]}\n字段：{row[3]}",
                f"{row[4]}\n目标格式：{row[11]}",
                f"现状：{row[5]}\n状态：{row[6]}\n优先级：{row[12]}",
                f"验收：{row[14]}\nETL/降级：{row[15]}",
            ])
        add_table(
            doc,
            ["ID", "数据组", "客户应提供的内容", "模型用途 / 建议格式", "当前状态", "验收与 ETL 后处理"],
            compact,
            [0.48, 1.0, 2.15, 1.9, 1.45, 3.0],
            font_size=7.1,
        )

    add_heading(doc, "5. 泵站数据能否直接给模型使用？", 1)
    add_callout(
        doc,
        "结论：大概率不能直接使用。",
        "客户给的泵站数据必须拆成静态资产、性能曲线/控制规则和运行时序三类，并通过 pump_id、curve_id、link_id 与 SWMM 网络对象建立稳定关系。只有完成这三类映射和质量检查后，才能编译为 SWMM 的 [PUMPS]、[CURVES]、[CONTROLS] 及必要的时间序列。",
        LIGHT_RED,
    )
    add_table(
        doc,
        ["客户可能提供的文件", "缺少时不能回答的问题", "ETL 必做动作", "模型输出"],
        [
            ["泵站/机组台账：泵 ID、站点、连接管段、型号、额定流量和扬程", "这台泵对应 SWMM 哪个 link？容量和单位是什么？", "统一 pump_id；映射 link_id；统一 m³/s、m；去重并绑定有效期", "SWMM [PUMPS] 资产记录"],
            ["泵曲线：Q-H、效率、并联/备用关系", "给定水位/扬程时泵能排多少？多台泵如何叠加？", "解析曲线点；检查单调性、单位和运行范围；生成 curve_id", "SWMM [CURVES]"],
            ["启停和控制逻辑：启泵水位、停泵水位、延迟、优先级、故障模式", "模型何时启动/停止？故障或手动模式如何表达？", "结构化为条件—动作规则；校验阈值和时间延迟；形成版本化控制表", "SWMM [CONTROLS] / scenario config"],
            ["SCADA / historian：时间、状态、转速、流量、扬程、水位、告警", "历史事件中泵实际是否运行、运行多久、实际排水多少？", "标签映射；转 UTC；重采样；状态/流量一致性检查；保留质量码", "泵运行时序、校准事件包"],
            ["仅有静态设施点或泵站名称", "不能确定实际排水能力和历史动作", "只能作为资产参考或情景位置，不能直接作为动态泵输入", "不进入工程泵控制；标记 MISSING / ASSUMED"],
        ],
        [2.1, 2.3, 3.0, 2.0],
        font_size=7.8,
    )

    add_heading(doc, "6. 客户字段级模板", 1)
    add_body(doc, "以下字段是最低交付模板。字段可以分散在多个客户表或图层中，但 ETL 后必须形成稳定的数据表名、字段名、类型、单位和跨表 ID 关系。动态数据统一使用 ISO 8601 UTC；缺失值必须使用 NULL + quality_code，不得用 0 代替。")
    template_groups: dict[str, list[list[str]]] = {}
    for row in TEMPLATES:
        template_groups.setdefault(row[0], []).append(row)
    template_labels = {
        "event_catalog": "事件目录",
        "rain_gauge": "雨量站时序",
        "radar_qpe": "雷达 QPE/QPF",
        "network_pipe": "管段",
        "network_node": "节点",
        "pump_asset": "泵资产",
        "pump_scada": "泵 SCADA",
        "outfall_boundary": "出水口边界",
        "flood_observation": "积水/水位观测",
        "network_change": "管网变更",
        "model_label": "模型派生标签（非客户原始输入）",
    }
    field_section = doc.add_section(start_type=1)
    set_landscape(field_section)
    for group, rows in template_groups.items():
        add_heading(doc, f"6.{list(template_groups).index(group) + 1} {template_labels.get(group, group)}", 2)
        add_table(
            doc,
            ["字段名", "中文名称", "类型", "单位 / 坐标", "必填", "示例", "校验规则", "说明"],
            [[r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8]] for r in rows],
            [1.25, 1.35, 0.9, 1.0, 0.55, 1.55, 2.45, 2.15],
            font_size=7.1,
        )

    add_heading(doc, "7. ETL 流程与模型原生输出", 1)
    add_table(
        doc,
        ["ETL 步骤", "输入", "处理", "输出", "不通过时"],
        [
            ["1. 接收登记", "文件、数据库、接口和图纸", "dataset_id、版本、责任方、许可、哈希", "landing manifest", "拒绝无来源/无版本数据"],
            ["2. 空间与高程统一", "GIS、栅格、测量和水位", "CRS、垂直基准、单位、NoData 和精度元数据", "标准化空间层", "隔离基准不明数据"],
            ["3. ID 与拓扑治理", "管线、节点、进水口、出水口、泵站", "唯一 ID、端点吸附、上下游、孤立边、重复对象、边界映射", "network_pipe/node/inlet/outfall", "不进入工程网络"],
            ["4. 参数标准化", "断面、粗糙、入渗、土地利用、曲线", "单位转换、代码域映射、范围检查、来源等级", "canonical parameter tables", "采用情景集合并标注假设"],
            ["5. 时间对齐", "降雨、潮位、SCADA、观测", "UTC、采样间隔、累计语义、重采样、缺测和质量码", "event package", "降级为代理或不进入校准"],
            ["6. 编译模型输入", "标准化表、栅格和参数", "生成 SWMM .inp、ANUGA mesh/quantity/boundary 和耦合映射", "model input bundle", "阻止运行或只允许诊断运行"],
            ["7. 运行质量门", "输入包和模型版本", "拓扑、连续性、质量守恒、文件哈希、回执", "run receipt + QA report", "结果不可声明工程有效"],
        ],
        [1.0, 2.0, 3.0, 2.2, 1.8],
        font_size=7.7,
    )

    add_heading(doc, "8. 推荐客户交付目录", 1)
    add_body(doc, "建议客户按数据类别分包，并在根目录提供 manifest.yaml 或 manifest.json。文件名中的日期和版本不能替代 manifest 中的 valid_from、valid_to、received_at 和 sha256。")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_after = Pt(8)
    code_run = code.add_run(
        "abu_dhabi_hydro_delivery/\n"
        "├── manifest.yaml\n"
        "├── 01_network/\n"
        "│   ├── network.gdb/ or network.gpkg\n"
        "│   ├── pipes.csv / nodes.csv / inlets.csv / outfalls.csv\n"
        "│   └── topology_map.csv\n"
        "├── 02_terrain/\n"
        "│   ├── dtm_5m.tif\n"
        "│   ├── terrain_metadata.json\n"
        "│   └── microtopography.gpkg / lidar.laz (if available)\n"
        "├── 03_surface_and_catchments/\n"
        "│   ├── catchments.gpkg / catchments.csv\n"
        "│   ├── landuse.tif or landuse.gpkg\n"
        "│   └── soil_infiltration.csv\n"
        "├── 04_forcing/\n"
        "│   ├── rainfall/\n"
        "│   ├── tide_outfall/\n"
        "│   └── pumps/ (assets, curves, controls, scada)\n"
        "├── 05_observations/\n"
        "│   ├── network_levels_flow.parquet\n"
        "│   └── inundation_observations.gpkg / photos_index.csv\n"
        "├── 06_events/\n"
        "│   ├── event_catalog.csv\n"
        "│   └── event_packages/\n"
        "└── 07_reference/\n"
        "    ├── data_dictionary.xlsx\n"
        "    └── drawings_and_reports/"
    )
    code_run.font.name = "Menlo"
    code_run.font.size = Pt(8.5)
    code_run.font.color.rgb = RGBColor.from_string(NAVY)

    add_heading(doc, "9. 质量验收最低规则", 1)
    quality_rows = [[r[0], r[1], r[2], r[3], r[4]] for r in QUALITY]
    add_table(doc, ["编号", "检查域", "验收规则", "建议阈值", "失败处理"], quality_rows, [0.65, 1.15, 4.3, 1.55, 2.2], font_size=7.6)

    add_heading(doc, "10. 客户 P0 优先交付清单", 1)
    p0_rows = []
    for row in PRIORITY:
        if row[3] == "P0":
            p0_rows.append([row[0], row[1], row[2], row[4], row[5], row[6]])
    add_table(doc, ["编号", "客户交付项", "服务模型", "建议责任方", "目的", "对应需求"], p0_rows, [0.65, 3.0, 1.25, 1.7, 2.5, 1.8], font_size=7.6)

    add_heading(doc, "11. 交付前最终检查清单", 1)
    for text in [
        "每个空间数据集都有 CRS、水平单位、垂直基准、数据日期和责任方。",
        "每个管线两端都能解析到有效节点；上下游、流向、断面和管底高程可追溯。",
        "每个泵站都能从泵台账映射到 SWMM link；泵曲线和控制规则有明确版本。",
        "每条动态时序都使用 UTC，说明采样/累计间隔，保留质量码和缺测标记。",
        "降雨、潮位、泵站、观测数据均能通过 event_id 与同一事件包关联。",
        "DTM、管网高程、潮位和观测水位使用同一垂直基准或提供转换关系。",
        "ETL 后生成的 SWMM .inp、ANUGA 网格/量和映射表都有版本、哈希和运行回执。",
        "任何默认值、代理值或自动派生值均有 provenance、影响等级和替换路径。",
    ]:
        add_bullet(doc, text)

    add_callout(
        doc,
        "对客户的最简要求",
        "客户不需要直接制作 SWMM .inp 或 ANUGA Python 对象；客户应提供有来源、有字段字典、有时间/空间基准的源数据包。ETL 团队负责把这些源数据转换成模型可运行的标准输入，并把每次转换和质量结果回写到 manifest 和 run receipt。",
        LIGHT_CYAN,
    )

    doc.save(DOCX_PATH)


def xlsx_style(ws, widths: Sequence[float]) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    thin = Side(style="thin", color=GRID)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Arial", size=9, color=TEXT)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 48
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def add_xlsx_sheet(wb: Workbook, name: str, headers: Sequence[str], rows: Iterable[Sequence[object]], widths: Sequence[float]) -> None:
    ws = wb.create_sheet(name)
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    xlsx_style(ws, widths)


def build_xlsx() -> None:
    wb = Workbook()
    overview = wb.active
    overview.title = "00_说明"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:F1")
    overview["A1"] = "阿布扎比一二维水动力模型客户数据输入与 ETL 规范"
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 34
    overview.merge_cells("A2:F2")
    overview["A2"] = "V1.0 · 2026-09-22 · 1D EPA SWMM + 2D ANUGA + coupling ETL"
    overview["A2"].font = Font(name="Arial", size=10, color=MUTED)
    overview.merge_cells("A4:F6")
    overview["A4"] = (
        "核心结论：客户原始数据不能直接作为模型输入。必须经过字段、单位、坐标、垂直基准、时间、ID、拓扑、质量和版本治理，"
        "再编译为 SWMM .inp、ANUGA 网格/量/边界和一二维交换映射。泵站数据必须拆为泵资产、泵曲线/控制和 SCADA 时序。"
    )
    overview["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    overview["A4"].fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    overview["A4"].font = Font(name="Arial", size=11, color=TEXT)
    for row in range(4, 7):
        overview.row_dimensions[row].height = 34
    overview["A8"] = "数据层级"
    overview["A8"].font = Font(name="Arial", size=12, bold=True, color=NAVY)
    overview.append(["层级", "对象", "示例格式", "是否直接运行", "关键 ETL", "说明"])
    overview_rows = [
        ["L0", "客户源数据", "GDB / GeoTIFF / CSV / Parquet / PDF", "否", "读取、登记、字段字典", "保留原始文件，不覆盖"],
        ["L1", "标准化数据层", "GeoPackage / PostGIS / Parquet / COG / NetCDF", "否", "ID、CRS、UTC、单位、拓扑、质量码", "用于检查和编译"],
        ["L2", "模型原生输入", ".inp / mesh / quantity / JSONL", "质量门后", "生成模型、哈希和 run receipt", "用于 SWMM / ANUGA 运行"],
    ]
    for row in overview_rows:
        overview.append(row)
    xlsx_style(overview, [14, 28, 30, 15, 38, 35])

    add_xlsx_sheet(
        wb,
        "01_需求矩阵",
        ["ID", "数据组", "数据项", "核心字段/内容", "模型作用", "当前基础", "满足情况", "期望来源", "空间要求", "时间要求", "历史覆盖", "建议格式", "优先级", "阶段", "验收标准", "缺失时 ETL/降级"],
        HYDRO,
        [9, 14, 28, 42, 35, 32, 13, 30, 24, 22, 22, 24, 10, 16, 48, 42],
    )
    add_xlsx_sheet(
        wb,
        "02_字段模板",
        ["数据表", "字段名", "中文名称", "类型", "单位/坐标", "必填", "示例", "校验规则", "说明"],
        TEMPLATES,
        [22, 26, 25, 16, 16, 12, 30, 42, 40],
    )
    add_xlsx_sheet(
        wb,
        "03_质量规则",
        ["规则编号", "检查域", "验收规则", "建议阈值", "失败处理"],
        QUALITY,
        [12, 18, 62, 28, 32],
    )
    add_xlsx_sheet(
        wb,
        "04_P0交付清单",
        ["编号", "交付项", "服务模型", "优先级", "责任方", "预期价值", "对应需求"],
        [row[:7] for row in PRIORITY if row[3] == "P0"],
        [10, 45, 20, 10, 26, 42, 30],
    )
    pump_rows = [
        ["pump_asset", "pump_id / station_id / link_id / curve_id / valid_from / valid_to", "静态资产与模型对象映射", "必须", "CSV / Parquet / GDB + 数据字典"],
        ["pump_curve", "curve_id / x / y / unit / curve_type", "Q-H 或其他性能曲线点", "必须", "CSV / 曲线文件 / 设备报告"],
        ["pump_control", "pump_id / condition / threshold / action / delay / priority", "启停、延迟、故障和备用规则", "工程模型必需", "CSV / JSON / 控制逻辑说明"],
        ["pump_scada", "timestamp_utc / pump_id / state / speed / flow_m3s / head_m / level / alarm / quality_code", "历史实际动作与能力", "校准事件必需", "CSV / Parquet / PI 导出"],
        ["mobile_pump_action", "event_id / timestamp_utc / location / capacity_m3s / head_m / duration_s / status", "移动泵应急情景，不是默认管网输入", "情景分析时", "CSV / JSON"],
    ]
    add_xlsx_sheet(wb, "05_泵站输入示例", ["数据表", "必备字段", "作用", "必填级别", "建议格式"], pump_rows, [22, 70, 35, 18, 32])
    wb.save(XLSX_PATH)


def write_json() -> None:
    payload = {
        "metadata": {
            "title": "阿布扎比一二维水动力模型客户数据输入与 ETL 规范",
            "version": "V1.0",
            "date": "2026-09-22",
            "scope": "EPA SWMM 1D, ANUGA 2D and 1D-2D coupling",
        },
        "hydro_requirements": HYDRO,
        "field_templates": TEMPLATES,
        "quality_rules": QUALITY,
        "priority_delivery": PRIORITY,
    }
    JSON_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def verify() -> dict[str, object]:
    assert DOCX_PATH.exists() and DOCX_PATH.stat().st_size > 10_000
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=False)
    assert wb.sheetnames == ["00_说明", "01_需求矩阵", "02_字段模板", "03_质量规则", "04_P0交付清单", "05_泵站输入示例"]
    return {
        "docx": str(DOCX_PATH),
        "xlsx": str(XLSX_PATH),
        "json": str(JSON_PATH),
        "hydro_requirements": len(HYDRO),
        "field_templates": len(TEMPLATES),
        "quality_rules": len(QUALITY),
    }


def main() -> None:
    build_docx()
    build_xlsx()
    write_json()
    print(json.dumps(verify(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
