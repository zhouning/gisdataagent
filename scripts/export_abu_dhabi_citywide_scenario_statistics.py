#!/usr/bin/env python3
"""Export auditable statistics for the registered citywide 2 h scenarios."""

from __future__ import annotations

import json
import math
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE = Path(
    "/Users/zhouning/.local/share/gisdataagent/private/abu_dhabi_stormwater/"
    "citywide_coupled_matrix_v3_land_masked_20260926"
)
OUTPUT = Path("/Users/zhouning/Downloads/阿布扎比全市2小时33-61-96mm_水动力统计指标.xlsx")
CELL_AREA_M2 = 250.0 * 250.0
THRESHOLDS_M = (0.01, 0.05, 0.15, 0.30, 0.50, 1.00)


def _root(mm: int) -> Path:
    return BASE / (
        f"abu-dhabi-citywide-coupled-{mm}mm-2h-250m-one-way-land-masked-74h/"
        "runs/full-74h-one-way-swmm-to-anuga-routing30s"
    )


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _percentile(values: list[float], percentile: float) -> float | None:
    if not values:
        return None
    values = sorted(values)
    position = (len(values) - 1) * percentile
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return values[lower]
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def _safe_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _depth_stats(features: list[dict], threshold: float) -> dict[str, float | int | None]:
    selected = [f for f in features if float((f.get("properties") or {}).get("maximum_depth_m") or 0.0) >= threshold]
    depths = [float((f.get("properties") or {}).get("maximum_depth_m") or 0.0) for f in selected]
    final_depths = [float((f.get("properties") or {}).get("final_depth_m") or 0.0) for f in selected]
    final_selected = [value for value in final_depths if value >= threshold]
    return {
        "count": len(selected),
        "area_m2": len(selected) * CELL_AREA_M2,
        "mean_depth_m": mean(depths) if depths else None,
        "median_depth_m": median(depths) if depths else None,
        "p95_depth_m": _percentile(depths, 0.95),
        "max_depth_m": max(depths) if depths else None,
        "final_count": len(final_selected),
        "final_area_m2": len(final_selected) * CELL_AREA_M2,
        "final_mean_depth_m": mean(final_depths) if final_depths else None,
        "final_max_depth_m": max(final_depths) if final_depths else None,
    }


def _receipt_stats(receipt: dict) -> dict[str, float | int | None]:
    windows = receipt.get("windows") or []
    sums = {}
    for field in (
        "rainfall_inflow_m3",
        "swmm_overflow_to_surface_m3",
        "head_exchange_swmm_to_surface_m3",
        "head_exchange_surface_to_swmm_m3",
        "total_swmm_to_anuga_m3",
        "total_anuga_to_swmm_m3",
        "external_outflow_m3",
        "infiltration_loss_m3",
    ):
        sums[field] = sum(float(row.get(field) or 0.0) for row in windows)
    max_surface = max((float(row.get("surface_storage_end_m3") or 0.0) for row in windows), default=0.0)
    max_network = max((float(row.get("swmm_node_storage_end_m3") or 0.0) for row in windows), default=0.0)
    max_residual = max((abs(float(row.get("coupled_mass_balance_residual_m3") or 0.0)) for row in windows), default=0.0)
    max_overflow_rate = max(
        (float(row.get("swmm_overflow_to_surface_m3") or 0.0) / 300.0 for row in windows),
        default=0.0,
    )
    max_external_rate = max(
        (float(row.get("external_outflow_m3") or 0.0) / 300.0 for row in windows),
        default=0.0,
    )
    return {
        **sums,
        "window_count": len(windows),
        "quality_pass_count": sum(bool(row.get("quality_passed")) for row in windows),
        "max_surface_storage_m3": max_surface,
        "max_network_storage_m3": max_network,
        "max_mass_balance_residual_m3": max_residual,
        "max_overflow_rate_m3s": max_overflow_rate,
        "max_external_outflow_rate_m3s": max_external_rate,
        "final_surface_storage_m3": float(windows[-1].get("surface_storage_end_m3") or 0.0) if windows else None,
        "final_network_storage_m3": float(windows[-1].get("swmm_node_storage_end_m3") or 0.0) if windows else None,
    }


def _scenario(mm: int) -> dict:
    root = _root(mm)
    summary = _read_json(root / "delivery_summary.json")
    maximum = _read_json(root / "maximum_depth_wgs84.geojson")
    volume = _read_json(root / "inundation_volume_points_wgs84.geojson")
    volume_summary = _read_json(root / "volume_summary.json")
    volume_ts = _read_json(root / "volume_timeseries.json")
    receipt = _read_json(root / "bidirectional_coupling_receipt.json")
    features = maximum.get("features") or []
    points = volume.get("features") or []
    frames = volume_ts.get("frames") or []
    domain = summary.get("domain") or {}
    delivery = summary.get("delivery") or {}
    results = summary.get("results") or {}
    forcing = summary.get("forcing") or {}
    coupling = summary.get("coupling") or {}
    surface_stats = volume_summary
    max_area_frame = max(frames, key=lambda row: int(row.get("inundated_cell_count_ge_threshold") or 0), default={})
    max_storage_frame = max(frames, key=lambda row: float(row.get("total_surface_water_volume_m3") or 0.0), default={})
    max_cell_volume = max(
        (float((f.get("properties") or {}).get("maximum_volume_m3") or 0.0) for f in points),
        default=0.0,
    )
    sum_cell_peak_volumes = sum(
        float((f.get("properties") or {}).get("maximum_volume_m3") or 0.0) for f in points
    )
    active_area_m2 = float(domain.get("active_land_cells") or 0) * CELL_AREA_M2
    rainfall_mm = float(forcing.get("total_depth_mm") or mm)
    rainfall_volume_estimate = rainfall_mm / 1000.0 * active_area_m2
    depth_rows = {threshold: _depth_stats(features, threshold) for threshold in THRESHOLDS_M}
    receipt_stats = _receipt_stats(receipt)
    header = {
        "scenario": f"{mm} mm / 2 h",
        "rainfall_mm": rainfall_mm,
        "rainfall_duration_h": float(forcing.get("duration_minutes") or 120.0) / 60.0,
        "model_type": "SWMM → ANUGA 2D 单向耦合",
        "status": summary.get("status"),
        "grid_m": float(domain.get("cell_size_m") or 250.0),
        "dtm_resolution_m": float((summary.get("surface") or {}).get("source_resolution_m", [5.0])[0]),
        "active_land_cells": int(domain.get("active_land_cells") or 0),
        "active_land_area_km2": active_area_m2 / 1_000_000.0,
        "rainfall_volume_estimate_m3": rainfall_volume_estimate,
        "max_envelope_area_ge_0_01_km2": depth_rows[0.01]["area_m2"] / 1_000_000.0,
        "max_simultaneous_area_ge_0_01_km2": float(max_area_frame.get("inundated_cell_count_ge_threshold") or 0) * CELL_AREA_M2 / 1_000_000.0,
        "max_simultaneous_area_time_min": float(max_area_frame.get("time_minutes") or 0.0),
        "max_depth_m": float(results.get("maximum_depth_m") or 0.0),
        "max_depth_time_min": max((float((f.get("properties") or {}).get("maximum_depth_time_minutes") or 0.0) for f in features), default=0.0),
        "final_max_depth_m": max((float((f.get("properties") or {}).get("final_depth_m") or 0.0) for f in features), default=0.0),
        "max_total_surface_storage_m3": float(max_storage_frame.get("total_surface_water_volume_m3") or 0.0),
        "max_total_surface_storage_time_min": float(max_storage_frame.get("time_minutes") or 0.0),
        "final_surface_storage_m3": float((frames[-1] if frames else {}).get("total_surface_water_volume_m3") or 0.0),
        "max_single_cell_peak_volume_m3": max_cell_volume,
        "sum_cell_peak_volumes_m3_not_simultaneous": sum_cell_peak_volumes,
        "swmm_overflow_to_surface_m3": receipt_stats["swmm_overflow_to_surface_m3"],
        "max_swmm_overflow_rate_m3s_equivalent": receipt_stats["max_overflow_rate_m3s"],
        "external_outflow_m3": receipt_stats["external_outflow_m3"],
        "max_external_outflow_rate_m3s_equivalent": receipt_stats["max_external_outflow_rate_m3s"],
        "coupled_max_network_storage_m3": receipt_stats["max_network_storage_m3"],
        "coupled_final_network_storage_m3": receipt_stats["final_network_storage_m3"],
        "head_exchange_back_to_swmm_m3": receipt_stats["total_anuga_to_swmm_m3"],
        "max_mass_balance_residual_m3": receipt_stats["max_mass_balance_residual_m3"],
        "coupling_windows": receipt_stats["window_count"],
        "expected_windows": int(delivery.get("expected_window_count") or 0),
        "complete_recession_verified": bool(delivery.get("complete_recession_verified", False)),
        "coupling_quality_passed": bool(coupling.get("quality_passed", False)),
        "source_root": str(root),
    }
    return {
        "mm": mm,
        "root": root,
        "summary": summary,
        "maximum": maximum,
        "volume": volume,
        "volume_summary": surface_stats,
        "frames": frames,
        "receipt": receipt,
        "receipt_stats": receipt_stats,
        "depth_rows": depth_rows,
        "header": header,
    }


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[object]], table_name: str | None = None) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value or "")) for cell in column_cells[:200])
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 34)
    if len(rows) > 0:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        # Excel table names are identifiers, not display labels; keep them ASCII
        # so Chinese worksheet titles do not cause compatibility issues.
        table = Table(displayName=table_name or f"T_{title.replace(' ', '_')}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    ws.sheet_view.showGridLines = False


def main() -> None:
    scenarios = [_scenario(mm) for mm in (33, 61, 96)]
    wb = Workbook()
    wb.remove(wb.active)

    summary_headers = [
        "场景", "降雨量_mm", "时长_h", "模型类型", "结果状态", "计算网格_m", "输入DTM_m",
        "有效陆域单元", "有效陆域面积_km²", "均匀降雨几何体积估算_m³",
        "最大包络积水面积≥0.01m_km²", "最大同一时刻积水面积≥0.01m_km²", "最大同一时刻面积时刻_min",
        "最大积水深度_m", "最大深度时刻_min", "末时间片最大深度_m",
        "最大地表瞬时蓄水量_m³", "蓄水峰值时刻_min", "末时间片地表蓄水量_m³",
        "最大单网格峰值体积_m³", "各网格峰值体积求和_m³_不可视为同时值",
        "SWMM→二维累计溢出源项_m³", "最大5分钟溢出源项_m³/s等效", "累计外排量_m³", "最大5分钟外排量_m³/s等效",
        "耦合回执管网节点蓄水峰值_m³", "耦合回执末时刻管网蓄水_m³", "二维→SWMM回流量_m³",
        "最大耦合质量残差_m³", "已完成窗口", "预期窗口", "完整退水已验证", "耦合质量门通过", "结果目录",
    ]
    _write_sheet(wb, "场景汇总", summary_headers, [[s["header"][key] for key in (
        "scenario", "rainfall_mm", "rainfall_duration_h", "model_type", "status", "grid_m", "dtm_resolution_m",
        "active_land_cells", "active_land_area_km2", "rainfall_volume_estimate_m3", "max_envelope_area_ge_0_01_km2",
        "max_simultaneous_area_ge_0_01_km2", "max_simultaneous_area_time_min", "max_depth_m", "max_depth_time_min",
        "final_max_depth_m", "max_total_surface_storage_m3", "max_total_surface_storage_time_min", "final_surface_storage_m3",
        "max_single_cell_peak_volume_m3", "sum_cell_peak_volumes_m3_not_simultaneous", "swmm_overflow_to_surface_m3",
        "max_swmm_overflow_rate_m3s_equivalent", "external_outflow_m3", "max_external_outflow_rate_m3s_equivalent",
        "coupled_max_network_storage_m3", "coupled_final_network_storage_m3", "head_exchange_back_to_swmm_m3",
        "max_mass_balance_residual_m3", "coupling_windows", "expected_windows", "complete_recession_verified",
        "coupling_quality_passed", "source_root")] for s in scenarios], table_name="T_Summary")

    depth_headers = [
        "场景", "阈值_m", "最大包络单元数", "最大包络面积_m²", "最大包络面积_km²", "有效陆域占比_%",
        "包络内最大深度均值_m", "包络内最大深度中位数_m", "包络内最大深度P95_m", "包络内最大深度最大值_m",
        "末时间片单元数", "末时间片面积_km²", "末时间片平均深度_m", "末时间片最大深度_m",
    ]
    depth_rows = []
    for s in scenarios:
        active_area = s["header"]["active_land_area_km2"]
        for threshold, stats in s["depth_rows"].items():
            depth_rows.append([
                s["header"]["scenario"], threshold, stats["count"], stats["area_m2"], stats["area_m2"] / 1_000_000.0,
                stats["area_m2"] / 1_000_000.0 / active_area * 100.0 if active_area else None,
                stats["mean_depth_m"], stats["median_depth_m"], stats["p95_depth_m"], stats["max_depth_m"],
                stats["final_count"], stats["final_area_m2"] / 1_000_000.0, stats["final_mean_depth_m"], stats["final_max_depth_m"],
            ])
    _write_sheet(wb, "深度面积分布", depth_headers, depth_rows, table_name="T_DepthArea")

    time_headers = ["场景", "降雨量_mm", "时间_min", "总地表水量_m³", "积水单元数≥0.01m", "同一时刻积水面积_km²", "完整退水已验证"]
    time_rows = []
    for s in scenarios:
        for frame in s["frames"]:
            cells = int(frame.get("inundated_cell_count_ge_threshold") or 0)
            time_rows.append([
                s["header"]["scenario"], s["header"]["rainfall_mm"], frame.get("time_minutes"),
                frame.get("total_surface_water_volume_m3"), cells, cells * CELL_AREA_M2 / 1_000_000.0,
                s["header"]["complete_recession_verified"],
            ])
    _write_sheet(wb, "地表水量时序", time_headers, time_rows, table_name="T_SurfaceTimeseries")

    coupling_headers = ["场景", "指标", "数值", "单位", "统计方式/定义", "来源"]
    coupling_rows = []
    coupling_metrics = [
        ("累计SWMM→二维溢出源项", "swmm_overflow_to_surface_m3", "m³", "887个已完成窗口逐窗求和", "bidirectional_coupling_receipt.json"),
        ("最大5分钟溢出源项等效流量", "max_overflow_rate_m3s", "m³/s", "单窗口溢出体积 / 300 s 的最大值", "bidirectional_coupling_receipt.json"),
        ("累计外排量", "external_outflow_m3", "m³", "887个已完成窗口逐窗求和", "bidirectional_coupling_receipt.json"),
        ("最大5分钟外排等效流量", "max_external_outflow_rate_m3s", "m³/s", "单窗口外排体积 / 300 s 的最大值", "bidirectional_coupling_receipt.json"),
        ("最大地表蓄水量（耦合回执）", "max_surface_storage_m3", "m³", "逐窗surface_storage_end_m3最大值", "bidirectional_coupling_receipt.json"),
        ("最大管网节点蓄水量（耦合回执）", "max_network_storage_m3", "m³", "逐窗swmm_node_storage_end_m3最大值；不是静态设计库容", "bidirectional_coupling_receipt.json"),
        ("二维→SWMM累计回流", "total_anuga_to_swmm_m3", "m³", "逐窗求和；本批次为0", "bidirectional_coupling_receipt.json"),
        ("最大耦合质量残差", "max_mass_balance_residual_m3", "m³", "逐窗绝对值最大值", "bidirectional_coupling_receipt.json"),
    ]
    for s in scenarios:
        rs = s["receipt_stats"]
        for label, key, unit, definition, source in coupling_metrics:
            coupling_rows.append([s["header"]["scenario"], label, rs.get(key), unit, definition, source])
        coupling_rows.extend([
            [s["header"]["scenario"], "已完成耦合窗口", rs["window_count"], "窗口", "实际存在的耦合回执窗口数", "bidirectional_coupling_receipt.json"],
            [s["header"]["scenario"], "质量门通过窗口", rs["quality_pass_count"], "窗口", "window.quality_passed为true的数量", "bidirectional_coupling_receipt.json"],
        ])
    _write_sheet(wb, "一维耦合统计", coupling_headers, coupling_rows, table_name="T_Coupling")

    point_headers = ["场景", "排名", "cell_id", "经度", "纬度", "最大瞬时体积_m³", "峰值水深_m", "平均水深_m", "峰值积水面积_m²", "峰值时刻_min", "单元最大深度_m", "末时间片体积_m³"]
    point_rows = []
    for s in scenarios:
        ordered = sorted(s["volume"].get("features") or [], key=lambda f: float((f.get("properties") or {}).get("maximum_volume_m3") or 0.0), reverse=True)[:50]
        for rank, feature in enumerate(ordered, 1):
            props = feature.get("properties") or {}
            coords = (feature.get("geometry") or {}).get("coordinates") or [None, None]
            point_rows.append([
                s["header"]["scenario"], rank, props.get("cell_id"), coords[0], coords[1],
                props.get("maximum_volume_m3"), props.get("depth_m_at_peak_volume"), props.get("mean_depth_m_at_peak_volume"),
                props.get("flooded_area_m2_at_peak_volume"), props.get("peak_volume_time_minutes"),
                props.get("maximum_depth_m"), props.get("final_volume_m3"),
            ])
    _write_sheet(wb, "积水点Top50", point_headers, point_rows, table_name="T_Top50")

    notes = [
        ["指标", "口径与限制"],
        ["最大包络积水面积", "在整个已有时间序列中曾达到≥阈值的250 m交付网格并集面积；不是同一时刻面积。"],
        ["最大同一时刻积水面积", "根据SWW逐帧三角形水深聚合，在单个5分钟输出帧内达到≥0.01 m的网格面积最大值。"],
        ["最大地表瞬时蓄水量", "逐帧将三角形水量聚合到网格后，取全市同一时刻总水量峰值；可代表本次场景的地表蓄水状态，不等于工程设计库容。"],
        ["最大单网格峰值体积", "每个250 m网格在其自身峰值时刻的最大瞬时体积，再取全市最大值。"],
        ["各网格峰值体积求和", "把每个网格的个人峰值相加，仅用于暴露空间峰值规模，不能解释为同一时刻全市水量。"],
        ["SWMM→二维累计溢出源项", "耦合回执中的swmm_overflow_to_surface_m3逐窗口求和，是SWMM向二维地表施加的溢出源项；不是独立的总降雨量。"],
        ["外排量", "耦合回执中的external_outflow_m3逐窗口求和。"],
        ["蓄存能力", "现有三场结果没有提供经批准的塘池/管网stage-area-volume曲线、设施容量或泵站容量，因此表中只报告事件内最大蓄水量，不声明设计蓄存能力。"],
        ["结果完整性", "三个场景均为887/888个窗口，最后5分钟缺失；完整退水未验证。所有末时间片指标只能称为已有窗口末值。"],
        ["模型类型", "本批次为客户5 m DTM输入、250 m ANUGA交付网格、SWMM→ANUGA单向耦合；没有动态水头反向反馈。"],
        ["数据来源", "delivery_summary.json、maximum_depth_wgs84.geojson、inundation_volume_points_wgs84.geojson、volume_timeseries.json、bidirectional_coupling_receipt.json。"],
    ]
    ws = wb.create_sheet("口径说明")
    for row in notes:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "scenario_count": len(scenarios), "time_rows": len(time_rows), "point_rows": len(point_rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
