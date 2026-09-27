from __future__ import annotations

from datetime import datetime
from pathlib import Path
import csv
import os

import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from data_agent.abu_dhabi_zone_b_design_storm import (
    ZONE_B_DDF_DEPTH_MM,
    ZONE_B_DURATIONS_MINUTES,
    ZONE_B_IDF_INTENSITY_MM_PER_HOUR,
    ZONE_B_SOURCE_YEAR,
    ZONE_B_TABLE_TITLE,
    ZONE_B_EVIDENCE_SHA256,
    zone_b_180_minute_hyetograph,
)


OUT = Path(os.environ.get(
    'ABU_DHABI_RAINFALL_PACKAGE_OUT',
    str(Path.cwd() / 'Zone_B官方DDF与雨型图_20260926'),
))
PRIMARY_SOURCE = r'D:\__!!_flood\IDD\Haydraulic Models-SWMP-IDS\Phase D Rev03\2. Collected Data\ADM-DMT\Rainfall Analysis Report _Phase B__V04.pdf'
CURVE_SOURCE = r'D:\__!!_flood\IDD\Haydraulic Models-SWMP-IDS\Phase D Rev03\2. Collected Data\ADM-DMT\Zone B Curve.docx'
MAIL_SOURCE = r'D:\__!!_flood\IDD\Haydraulic Models-SWMP-IDS\Phase D Rev03\2. Collected Data\ADM-DMT\FW_ Hyetograph of Abu Dhabi City - Updating SWMP for Abu Dhabi City.msg'
PRIMARY_SOURCE_DUP = r'D:\__!!_flood\IDD\Haydraulic Models-SWMP-IDS\Phase D Rev03\3. USWMP Phases\PHASE A Data Collection Report\Appendices\Appendix B - Data Collection\Appendix B.2 Collected Data\ADM-DMT\Rainfall Analysis Report _Phase B__V04.pdf'
CATALOG_SOURCE = os.environ.get(
    'ABU_DHABI_RAINFALL_CATALOG_SOURCE',
    'catalog_tree.html；雨水档案_目录与文件含义分析表.xlsx',
)
EVIDENCE_NOTE = '当前只有目录证据，原始 PDF/DOCX/MSG 未取回核验；表号、原始文档标题和页码需用原件最终确认。'


def choose_font() -> str:
    for candidate in ('Hiragino Sans GB', 'Songti SC', 'Heiti SC'):
        if any(candidate in f.name for f in font_manager.fontManager.ttflist):
            return candidate
    return 'DejaVu Sans'


def derived_hyetographs() -> dict[int, list[dict[str, float]]]:
    result: dict[int, list[dict[str, float]]] = {}
    for rp in sorted(ZONE_B_DDF_DEPTH_MM):
        series, _stats = zone_b_180_minute_hyetograph(
            rp, start=datetime(2000, 1, 1), peak_position_percent=40
        )
        rows = []
        for index, (_stamp, intensity) in enumerate(series[:-1]):
            rows.append({
                'index': index + 1,
                'elapsed_minute': index * 5,
                'depth_mm_per_5min': intensity / 12,
                'intensity_mm_per_hour': intensity,
            })
        result[rp] = rows
    return result


def style_sheet(ws, widths: dict[str, float] | None = None) -> None:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0B6E7F')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for key, width in (widths or {}).items():
        ws.column_dimensions[key].width = width


def make_workbook(hyetographs: dict[int, list[dict[str, float]]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = '说明与证据边界'
    rows = [
        ('主题', '阿布扎比 Zone B DDF 与派生雨型交付包'),
        ('官方表名', ZONE_B_TABLE_TITLE),
        ('资料年份', ZONE_B_SOURCE_YEAR),
        ('最可能主报告路径', PRIMARY_SOURCE),
        ('Phase A重复归档路径', PRIMARY_SOURCE_DUP),
        ('配套曲线文件', CURVE_SOURCE),
        ('配套邮件文件', MAIL_SOURCE),
        ('目录依据', CATALOG_SOURCE),
        ('证据SHA-256', ZONE_B_EVIDENCE_SHA256),
        ('已确认内容', '2/5/10/25/50/100年一遇、5/10/15/30/60/120/180/360/720/1440分钟累计雨量与平均雨强。'),
        ('未确认内容', '官方是否另有固定完整hyetograph、峰值位置、5分钟分配方法和原始页码。'),
        ('本包派生方法', '按官方DDF累计深度；180分钟内对数历时-深度插值到5分钟；递增雨量采用交替块法；峰值位置40%。'),
        ('使用边界', 'DDF数值可作为官方输入证据；5分钟曲线是模型派生方案，不应标注为官方观测雨型。'),
        ('核验状态', EVIDENCE_NOTE),
    ]
    ws.append(['字段', '内容'])
    for row in rows:
        ws.append(row)
    style_sheet(ws, {'A': 24, 'B': 120})
    ws.freeze_panes = 'A2'

    ws = wb.create_sheet('官方DDF累计雨量_mm')
    ws.append(['历时_min', *[f'{rp}年一遇' for rp in sorted(ZONE_B_DDF_DEPTH_MM)]])
    for index, duration in enumerate(ZONE_B_DURATIONS_MINUTES):
        ws.append([duration, *[ZONE_B_DDF_DEPTH_MM[rp][index] for rp in sorted(ZONE_B_DDF_DEPTH_MM)]])
    style_sheet(ws, {'A': 16, **{get_column_letter(i): 16 for i in range(2, 8)}})

    ws = wb.create_sheet('官方IDF平均雨强_mmh')
    ws.append(['历时_min', *[f'{rp}年一遇' for rp in sorted(ZONE_B_IDF_INTENSITY_MM_PER_HOUR)]])
    for index, duration in enumerate(ZONE_B_DURATIONS_MINUTES):
        ws.append([duration, *[ZONE_B_IDF_INTENSITY_MM_PER_HOUR[rp][index] for rp in sorted(ZONE_B_IDF_INTENSITY_MM_PER_HOUR)]])
    style_sheet(ws, {'A': 16, **{get_column_letter(i): 16 for i in range(2, 8)}})

    ws = wb.create_sheet('180min派生时程汇总')
    ws.append(['elapsed_minute', *[f'{rp}年一遇_depth_mm_5min' for rp in sorted(hyetographs)], *[f'{rp}年一遇_intensity_mmh' for rp in sorted(hyetographs)]])
    for i in range(36):
        ws.append([
            i * 5,
            *[hyetographs[rp][i]['depth_mm_per_5min'] for rp in sorted(hyetographs)],
            *[hyetographs[rp][i]['intensity_mm_per_hour'] for rp in sorted(hyetographs)],
        ])
    style_sheet(ws, {'A': 18, **{get_column_letter(i): 20 for i in range(2, 15)}})

    for rp in sorted(hyetographs):
        ws = wb.create_sheet(f'{rp}yr_180min_5min')
        ws.append(['序号', '累计时间_min', '5分钟雨深_mm', '雨强_mm/h', '说明'])
        for row in hyetographs[rp]:
            ws.append([row['index'], row['elapsed_minute'], row['depth_mm_per_5min'], row['intensity_mm_per_hour'], '模型派生；非官方逐5分钟观测'])
        style_sheet(ws, {'A': 10, 'B': 18, 'C': 18, 'D': 18, 'E': 32})

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
    wb.save(OUT / '阿布扎比Zone_B官方DDF与派生雨型.xlsx')

    with (OUT / 'Zone_B_180分钟_5分钟派生雨型_全部重现期.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.writer(handle)
        writer.writerow(['return_period_years', 'elapsed_minutes', 'depth_mm_per_5min', 'intensity_mm_per_hour', 'peak_position_assumption_percent', 'temporal_method'])
        for rp in sorted(hyetographs):
            for row in hyetographs[rp]:
                writer.writerow([rp, row['elapsed_minute'], row['depth_mm_per_5min'], row['intensity_mm_per_hour'], 40, 'log-duration/log-depth interpolation + alternating block'])


def make_figures(hyetographs: dict[int, list[dict[str, float]]]) -> None:
    plt.rcParams['font.family'] = choose_font()
    plt.rcParams['axes.unicode_minus'] = False
    rps = sorted(ZONE_B_DDF_DEPTH_MM)
    colors = plt.cm.viridis([0.12 + 0.76 * i / (len(rps) - 1) for i in range(len(rps))])

    fig, axes = plt.subplots(2, 1, figsize=(13, 10), constrained_layout=True)
    ax = axes[0]
    for color, rp in zip(colors, rps):
        ax.plot(ZONE_B_DURATIONS_MINUTES, ZONE_B_DDF_DEPTH_MM[rp], marker='o', color=color, label=f'{rp}-year')
    ax.set_xscale('log')
    ax.set_xlabel('Duration (minutes)')
    ax.set_ylabel('Cumulative depth (mm)')
    ax.set_title('Zone B official DDF values registered in project')
    ax.grid(True, alpha=.25)
    ax.legend(ncol=3, fontsize=9)

    ax = axes[1]
    for color, rp in zip(colors, rps):
        x = [row['elapsed_minute'] for row in hyetographs[rp]]
        y = [row['intensity_mm_per_hour'] for row in hyetographs[rp]]
        ax.plot(x, y, linewidth=2, color=color, label=f'{rp}-year')
    ax.axvline(40 / 100 * 180, color='black', linestyle='--', linewidth=1, alpha=.55, label='assumed peak position 40%')
    ax.set_xlabel('Elapsed time (minutes)')
    ax.set_ylabel('Intensity (mm/h)')
    ax.set_title('Derived 180-minute / 5-minute alternating-block hyetographs')
    ax.grid(True, alpha=.25)
    ax.legend(ncol=3, fontsize=9)
    fig.suptitle('Abu Dhabi Zone B: official DDF evidence + model-derived temporal pattern', fontsize=15, fontweight='bold')
    fig.savefig(OUT / 'Zone_B_DDF与派生雨型总览.png', dpi=220, bbox_inches='tight')
    plt.close(fig)

    rp = 10
    fig, ax1 = plt.subplots(figsize=(13, 6.5), constrained_layout=True)
    rows = hyetographs[rp]
    x = [row['elapsed_minute'] for row in rows]
    depth = [row['depth_mm_per_5min'] for row in rows]
    intensity = [row['intensity_mm_per_hour'] for row in rows]
    bars = ax1.bar(x, depth, width=4.2, color='#0284C7', alpha=.88, label='5-minute depth (mm)')
    ax1.set_xlabel('Elapsed time (minutes)')
    ax1.set_ylabel('Depth per 5 minutes (mm)', color='#0284C7')
    ax1.tick_params(axis='y', labelcolor='#0284C7')
    ax1.grid(axis='y', alpha=.22)
    ax2 = ax1.twinx()
    ax2.plot(x, intensity, color='#DC2626', linewidth=2.2, marker='o', markersize=3, label='Intensity (mm/h)')
    ax2.set_ylabel('Intensity (mm/h)', color='#DC2626')
    ax2.tick_params(axis='y', labelcolor='#DC2626')
    ax1.axvline(72, color='black', linestyle='--', linewidth=1, alpha=.6)
    ax1.text(74, max(depth) * .92, 'assumed peak position 40%\n(t = 72 min)', fontsize=9, va='top')
    ax1.set_title('Zone B 10-year / 180-minute / 5-minute derived hyetograph\nOfficial total depth = 28.71 mm; temporal distribution is a model assumption')
    handles = [bars, ax2.lines[0]]
    ax1.legend(handles, ['5-minute depth (mm)', 'Intensity (mm/h)'], loc='upper right')
    fig.savefig(OUT / 'Zone_B_10年一遇_180分钟_5分钟派生雨型.png', dpi=220, bbox_inches='tight')
    plt.close(fig)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    hyetographs = derived_hyetographs()
    make_workbook(hyetographs)
    make_figures(hyetographs)
    readme = f'''# Zone B DDF 与派生雨型交付包

## 原始档案目录路径

- 主报告（当前最可能包含 Table 3-6 的文件）：`{PRIMARY_SOURCE}`
- 配套曲线文件：`{CURVE_SOURCE}`
- 配套邮件：`{MAIL_SOURCE}`
- Phase A 重复归档：`{PRIMARY_SOURCE_DUP}`

目录索引依据：`{CATALOG_SOURCE}`。

## 重要证据边界

- 已登记内容：{ZONE_B_TABLE_TITLE}，2/5/10/25/50/100 年一遇、10个历时的DDF/IDF数值。
- 当前未完成：原始PDF/DOCX/MSG没有在本机取回，Table 3-6 的原始页码和其是否另含固定完整hyetograph尚未核验。
- 本包中的5分钟雨型：使用官方DDF累计深度，180分钟内对数插值到5分钟，再以交替块法排列，峰值位置假设为40%。它是模型派生输入，不是官方逐5分钟观测。
- 证据哈希（当前工程登记）：`{ZONE_B_EVIDENCE_SHA256}`。

## 文件

- `阿布扎比Zone_B官方DDF与派生雨型.xlsx`：说明、来源、官方DDF/IDF表、6套重现期派生时程。
- `Zone_B_DDF与派生雨型总览.png`：DDF曲线 + 6套派生时程。
- `Zone_B_10年一遇_180分钟_5分钟派生雨型.png`：10年一遇详细柱状/折线图。
'''
    (OUT / 'README_来源与使用边界.md').write_text(readme, encoding='utf-8')
    print(OUT)


if __name__ == '__main__':
    main()
