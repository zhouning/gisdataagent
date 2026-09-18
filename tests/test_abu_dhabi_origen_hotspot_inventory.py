from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import openpyxl
from pyproj import Transformer
from starlette.applications import Starlette
from starlette.testclient import TestClient

from data_agent.abu_dhabi_flood_delivery_report import phase5_report_html
from data_agent.api import abu_dhabi_flood_routes as flood_routes
from data_agent.uwm.abu_dhabi_flood.hotspot_inventory import (
    augment_private_bundle_with_static_prior,
    build_private_bundle,
)

HEADERS = [
    "#",
    "Hotspot_Area",
    "Hotspot_Location",
    "Latitude",
    "Longitude",
    "Criticality",
    "Developer_Private_Area",
    "Developer Name",
    "Network_Available",
    "Root_Cause",
    "Root_Cause_Details",
    "IDD Interventoin Planned",
    "Design Solution",
]


def _write_sheet(sheet, header_row: int, rows: list[list[object]], *, offset: int = 0) -> None:
    for column, header in enumerate(HEADERS, start=1 + offset):
        sheet.cell(row=header_row, column=column, value=header)
    for row_index, row in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(row, start=1 + offset):
            sheet.cell(row=row_index, column=column, value=value)


def _write_workbooks(root: Path) -> None:
    adm = openpyxl.Workbook()
    adm.remove(adm.active)
    current = adm.create_sheet("ADM List")
    _write_sheet(
        current,
        2,
        [
            [
                209,
                "City",
                "Tunnel",
                24.45,
                54.45,
                "Very High",
                "No",
                "N/A",
                "Yes",
                "The network needs an upgrade",
                "Upgrade",
                "Yes",
                "New inlet",
            ],
            [
                210,
                "City",
                "Street",
                24.46,
                54.46,
                "Medium   ",
                "No",
                "N/A",
                "No",
                "Absense of Drainage Network",
                "No pipes",
                "Yes",
                "Drainage pipe",
            ],
        ],
    )
    history = adm.create_sheet("Sheet2")
    _write_sheet(
        history,
        4,
        [
            [
                1,
                "City",
                "Old tunnel",
                24.451,
                54.451,
                "High",
                "No",
                "N/A",
                "Yes",
                "Capacity",
                "Old",
                "No",
                None,
            ]
        ],
        offset=1,
    )
    adm.create_sheet("Exec Summary").sheet_state = "hidden"
    adm.save(root / "20260909 ADM Stormwater Hotspot Analysis v4.4.xlsx")

    for municipality, name, latitude, longitude, criticality in (
        (
            "AAM",
            "20260909_AAM_Stormwater_Hotspot_Analysis_v4.2.xlsx",
            " 24.440000°",
            "54.440000°",
            "High",
        ),
        ("DRM", "20260909_DRM_Stormwater_Hotspot_Analysis_v4.2.xlsx", 24.47, 54.47, "Low"),
    ):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Analysis"
        _write_sheet(
            sheet,
            3,
            [
                [
                    1,
                    municipality,
                    f"{municipality} location",
                    latitude,
                    longitude,
                    criticality,
                    "No",
                    "N/A",
                    "Yes",
                    "Inadequate design capacity",
                    "Capacity",
                    "Yes",
                    "Upgrade network",
                ]
            ],
        )
        workbook.save(root / name)


def _write_depth_result(root: Path) -> Path:
    root.mkdir()
    polygon = [
        [54.449, 24.449],
        [54.451, 24.449],
        [54.451, 24.451],
        [54.449, 24.451],
        [54.449, 24.449],
    ]
    depth = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature",
                "geometry": {"type": "Polygon", "coordinates": [polygon]},
                "properties": {"maximum_depth_m": 0.06},
            }
        ],
    }
    depth_path = root / "maximum_depth_wgs84.geojson"
    depth_path.write_text(json.dumps(depth), encoding="utf-8")
    transformer = Transformer.from_crs("EPSG:4326", "EPSG:32640", always_xy=True)
    lower = transformer.transform(54.40, 24.40)
    upper = transformer.transform(54.50, 24.50)
    (root / "delivery_summary.json").write_text(
        json.dumps({"domain": {"bounds_epsg32640": [lower[0], lower[1], upper[0], upper[1]]}}),
        encoding="utf-8",
    )
    return depth_path


def _write_nodes(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "geometry": {"type": "Point", "coordinates": [54.4501, 24.4501]},
                        "properties": {"node_id": "n-test"},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )


def _write_model_grid(path: Path) -> None:
    transformer = Transformer.from_crs("EPSG:4326", "EPSG:32640", always_xy=True)
    center_x, center_y = transformer.transform(54.45, 24.45)
    x = np.arange(center_x - 500.0, center_x + 750.0, 250.0)
    y = np.arange(center_y - 500.0, center_y + 750.0, 250.0)
    np.savez_compressed(
        path,
        x=x,
        y=y,
        values=np.zeros((len(y), len(x)), dtype=np.float32),
        land_mask=np.ones((len(y) - 1, len(x) - 1), dtype=bool),
    )


def test_build_private_bundle_separates_history_and_preserves_claim_boundary(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source"
    output = tmp_path / "private"
    source.mkdir()
    _write_workbooks(source)
    depth_path = _write_depth_result(tmp_path / "rp005")
    node_path = tmp_path / "nodes.geojson"
    _write_nodes(node_path)
    grid_path = tmp_path / "terrain_grid.npz"
    _write_model_grid(grid_path)

    manifest = build_private_bundle(
        source,
        output,
        depth_results={5: depth_path},
        node_path=node_path,
        maximum_node_distance_m=100.0,
        model_grid_path=grid_path,
    )

    assert manifest["inventory"]["current_count"] == 4
    assert manifest["inventory"]["history_count"] == 1
    assert manifest["inventory"]["current_by_municipality"] == {"AAM": 1, "ADM": 2, "DRM": 1}
    assert manifest["privacy"]["raw_workbooks_copied"] is False
    assert manifest["gwm_admission"]["forbidden_roles"] == [
        "event_water_depth_label",
        "event_flood_extent_ground_truth",
    ]
    assert not list(output.glob("*.xlsx"))
    assert (output / "hotspots.parquet").is_file()
    assert manifest["gwm_static_prior"]["status"] == "ready_prospective_training_only"
    assert manifest["gwm_static_prior"]["feature_count"] == 15
    assert manifest["gwm_static_prior"]["active_feature_count"] > 0
    assert (output / "gwm_static_prior_250m.npz").is_file()
    with np.load(output / "gwm_static_prior_250m.npz") as prior:
        assert prior["features"].shape == (15, 4, 4)
        assert "origen_criticality_weighted_proximity_exp_1km" in set(
            prior["feature_names"].tolist()
        )
        assert float(prior["features"].max()) > 0.0
    prior_receipt = json.loads(
        (output / "gwm_static_prior_receipt.json").read_text(encoding="utf-8")
    )
    assert prior_receipt["grid"]["source_sha256"]
    assert prior_receipt["source_record_counts"]["current_contributing_within_cutoff"] > 0
    assert prior_receipt["admission"]["forbidden"][0] == "retrofit_into_existing_frozen_models"

    quality = json.loads((output / "data_quality_receipt.json").read_text(encoding="utf-8"))
    assert quality["cached_error_value_count_workbook_wide"] == 0
    assert quality["error_value_samples"] == []

    current = json.loads((output / "hotspots_current.geojson").read_text(encoding="utf-8"))
    assert len(current["features"]) == 4
    aam = next(
        feature for feature in current["features"] if feature["properties"]["municipality"] == "AAM"
    )
    assert aam["geometry"]["coordinates"] == [54.44, 24.44]
    assert "latitude_degree_symbol_removed" in aam["properties"]["coordinate_normalizations"]
    first = next(feature for feature in current["features"] if feature["id"].startswith("ADM:209:"))
    assert first["properties"]["swmm_node_candidate"]["node_id"] == "n-test"
    assert first["properties"]["depth_concordance"]["rp005"]["hit_ge_0_05m"] is True

    concordance = json.loads(
        (output / "hotspot_spatial_concordance.json").read_text(encoding="utf-8")
    )
    assert concordance["status"] == "completed_weak_static_evidence"
    assert concordance["return_periods"]["5"]["hit_ge_0_05m_count"] == 1
    assert concordance["validation_gate"]["event_water_depth_validation"] == "pending"


def test_hotspot_routes_require_authentication_and_never_return_raw_rows(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "source"
    output = tmp_path / "private"
    source.mkdir()
    _write_workbooks(source)
    build_private_bundle(source, output)
    monkeypatch.setenv("ABU_DHABI_HOTSPOT_BUNDLE_ROOT", str(output))
    app = Starlette(routes=flood_routes.get_abu_dhabi_flood_routes())

    with TestClient(app) as client:
        assert client.get("/api/abu-dhabi/flood/hotspots/catalog").status_code == 401

    monkeypatch.setattr(
        flood_routes,
        "_get_user_from_request",
        lambda request: SimpleNamespace(identifier="analyst", metadata={"role": "analyst"}),
    )
    monkeypatch.setattr(flood_routes, "_set_user_context", lambda user: None)
    app = Starlette(routes=flood_routes.get_abu_dhabi_flood_routes())
    with TestClient(app) as client:
        catalog = client.get("/api/abu-dhabi/flood/hotspots/catalog")
        assert catalog.status_code == 200
        assert catalog.json()["inventory"]["current_count"] == 4
        response = client.get("/api/abu-dhabi/flood/hotspots/map?inventory=current")
        assert response.status_code == 200
        text = response.text
        assert '"raw"' not in text
        assert "source_formulas" not in text
        assert len(response.json()["features"]) == 4


def test_existing_private_bundle_can_be_augmented_without_rebuilding_evidence(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source"
    output = tmp_path / "private"
    source.mkdir()
    _write_workbooks(source)
    original = build_private_bundle(source, output)
    original_geojson_hash = original["artifacts"]["hotspots_current.geojson"]["sha256"]
    grid_path = tmp_path / "terrain_grid.npz"
    _write_model_grid(grid_path)

    augmented = augment_private_bundle_with_static_prior(output, grid_path)

    assert augmented["artifacts"]["hotspots_current.geojson"]["sha256"] == original_geojson_hash
    assert augmented["gwm_static_prior"]["feature_count"] == 15
    assert augmented["gwm_admission"]["frozen_confirmatory_model_use"] == "forbidden"


def test_phase5_report_defaults_to_english_and_contains_hotspot_evidence(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "source"
    output = tmp_path / "private"
    source.mkdir()
    _write_workbooks(source)
    build_private_bundle(source, output)
    monkeypatch.setenv("ABU_DHABI_HOTSPOT_BUNDLE_ROOT", str(output))

    html = phase5_report_html("en-US")

    assert "Phase 5 Delivery Report" in html
    assert "Current hotspots" in html
    assert "External GWM validation" in html
    assert "Prospective Origen GWM ablation" in html
    assert "NOT ADMITTED" in html
    assert not any("\u3400" <= character <= "\u9fff" for character in html)
